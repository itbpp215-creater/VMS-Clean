import re
import os
import json
import base64
import urllib.request
import urllib.error
import logging
from typing import Dict, Any, Optional, List
from pathlib import Path
from PIL import Image
import io
try:
    import pytesseract
except ImportError:
    pytesseract = None

try:
    import docx
except ImportError:
    docx = None

try:
    from pdf2image import convert_from_bytes, pdfinfo_from_bytes
except ImportError:
    convert_from_bytes = None
    pdfinfo_from_bytes = None

logger = logging.getLogger("nas-bot.ai_ocr")

def convert_thai_numerals(text: str) -> str:
    thai_to_arabic = str.maketrans('๐๑๒๓๔๕๖๗๘๙', '0123456789')
    return text.translate(thai_to_arabic)


def validate_thai_id_checksum(cid: Optional[str]) -> bool:
    """Validate Thai National ID card 13-digit checksum (Modulo 11)."""
    if not cid:
        return False
    digits = re.sub(r'\D', '', str(cid))
    if len(digits) != 13:
        return False
    if digits == "0000000000000" or digits.startswith("9"):
        return False
    try:
        total = sum(int(digits[i]) * (13 - i) for i in range(12))
        check_digit = (11 - (total % 11)) % 10
        return check_digit == int(digits[12])
    except (ValueError, IndexError):
        return False


def repair_thai_id_checksum(cid: Optional[str]) -> Optional[str]:
    """If a 13-digit Thai National ID fails Modulo-11 checksum due to Thai numeral OCR confusion (e.g. 4 vs 6), repair it."""
    if not cid:
        return None
    digits = re.sub(r'\D', '', str(cid))
    if len(digits) != 13:
        return digits
    if validate_thai_id_checksum(digits):
        return digits

    # Try specific Thai numeral confusion pairs on sequence digits (index 8 to 11): 4 <-> 6, 4 <-> 5, 5 <-> 6
    for idx in [10, 11, 9, 8]:
        orig = digits[idx]
        for alt in (['4', '5'] if orig == '6' else (['6', '5'] if orig == '4' else (['4', '6'] if orig == '5' else []))):
            cand = digits[:idx] + alt + digits[idx+1:]
            if validate_thai_id_checksum(cand):
                logger.info(f"Repaired Thai ID checksum ({digits} -> {cand}) at index {idx} ({orig}->{alt})")
                return cand
    return digits


def normalize_thai_name_candidate(raw_name: Optional[str]) -> Optional[str]:
    """Normalize OCR prefixes, correct common typos, and strip stray punctuation."""
    if not raw_name or not isinstance(raw_name, str):
        return None
    name = raw_name.strip(" \t\n\r\"'”’“‘()[]{}:;,./-")
    # Fix common OCR prefix misreadings
    name = re.sub(r'^(?:หาาย|บาย|ทาย|มาย|ผาย)\s*', 'นาย', name)
    name = re.sub(r'^(?:มางสาว|บา่งสาว)\s*', 'นางสาว', name)
    # Common OCR typo corrections & Transliterated/Half-Blood Names
    name = name.replace('พืรพัฒน์', 'พีรพัฒน์').replace('เมซี', 'เมฆี')
    name = name.replace('อุ่หหวงหิน', 'อุ่นหนองหิน').replace('อุนหนองหิน', 'อุ่นหนองหิน')
    name = name.replace('เอลิบคักดิ์', 'เฉลิมศักดิ์').replace('โกรษร', 'ไกรษร')
    name = name.replace('อีรศักดิ์', 'ธีรศักดิ์')
    # Half-blood transliterated corrections (e.g. Ludwig Schopf/Schopt)
    name = name.replace('ลุตวิก', 'ลุดวิก')
    name = name.replace('ซอบท์', 'ชอบท์').replace('ขอบท์', 'ชอบท์')
    name = name.replace('ชอบท่', 'ชอบท์').replace('ชอบท', 'ชอบท์')
    # Deduplicate consecutive karant characters
    name = re.sub(r'์+', '์', name)
    # Strip trailing punctuation/brackets but preserve Thai diacritics including karant
    name = re.sub(r'[\)\]\}\s]+$', '', name).strip()

    # Reject police unit, military, and administrative signboards mistakenly extracted as names
    POLICE_SIGNBOARD_KEYWORDS = [
        'ตชด', 'ชปส', 'กอ.รมน', 'ฉก.ทพ', 'อส.อ', 'ภ.จว', 'ศป.ปส', 'ร้อย.อส',
        'ร้อย ตชด', 'ร้อยอส', 'สภ.นางรอง', 'กองร้อย', 'ชป.ปส', 'สขว.', 'สภ.'
    ]
    if any(k in name for k in POLICE_SIGNBOARD_KEYWORDS):
        return None
    # Reject known phantom names parsed from unit signs
    if name in ['นางรี ร้อย', 'รี ร้อย', 'นางรอง', 'นายร้อย']:
        return None

    return name if name else None

def _remove_red_markup(img: Image.Image) -> Image.Image:
    """Filter out high-red markup, police stamps, and red watermarks for cleaner OCR."""
    img_rgb = img.convert('RGB')
    pixels = img_rgb.load()
    w, h = img_rgb.size
    for y in range(h):
        for x in range(w):
            r, g, b = pixels[x, y]
            # Aggressively wipe police red ink stamps, watermarks, and red highlights
            if (r > 110 and r > g * 1.25 and r > b * 1.25) or (r > 140 and g < 100 and b < 100):
                pixels[x, y] = (255, 255, 255)
    return img_rgb


def is_gibberish_text(text: str) -> bool:
    """Detect if text extracted from PDF is corrupted/unreadable font glyph encoding."""
    if not text or len(text.strip()) < 10:
        return False
    thai_chars = sum(1 for c in text if '\u0e00' <= c <= '\u0e7f')
    strange_chars = sum(1 for c in text if c in r'$^~{}[]\|`_#*')
    # If almost no Thai characters and multiple strange characters, or high strange ratio
    if thai_chars < 5 and strange_chars >= 3:
        return True
    if len(text) > 20 and (strange_chars / len(text)) > 0.05:
        return True
    return False


def extract_phone_list_data(image_bytes: bytes, text: Optional[str] = None) -> List[Dict[str, Any]]:
    """
    Extracts structured phone records with area codes/tags from a phone list image.
    e.g. '1. 0842674867 บร2', '10. 0807579072 บร1'
    """
    try:
        if not text:
            raw_img = Image.open(io.BytesIO(image_bytes))
            clean_img = _remove_red_markup(raw_img)
            text = pytesseract.image_to_string(clean_img, lang="tha+eng")

        entries = []
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        for line in lines:
            m_phone = re.search(r'\b(0[689]\d{8})\b|\b(0[689]\d[-\s]?\d{3}[-\s]?\d{4})\b', line)
            if m_phone:
                clean_phone = re.sub(r'\D', '', m_phone.group(0))
                tag_match = re.search(r'(บร\s*\d+|สต\s*\d+|พยัค\s*\d+แขบ|แขบ\s*\d+)', line)
                tag = tag_match.group(1).replace(' ', '') if tag_match else None
                
                idx_match = re.search(r'^(\d+)\.', line)
                idx = int(idx_match.group(1)) if idx_match else len(entries) + 1
                
                sheet_title = f"โพยชุด {tag}" if tag else "โพยชุดเบอร์โทร"
                entries.append({
                    "index": idx,
                    "phone_number": clean_phone,
                    "sheet_tag": tag,
                    "sheet_name": sheet_title,
                    "area_tag": tag,
                    "role_tag": "เบอร์ในโพยเครือข่าย",
                    "notes": f"สมุดจดเบอร์เครือข่าย ลำดับที่ {idx}: {clean_phone} {tag or ''}".strip(),
                    "raw_line": line
                })
        return entries
    except Exception as e:
        logger.error(f"Phone list extraction failed: {e}")
        return []


THAI_PROVINCES = [
    "กรุงเทพมหานคร", "กรุงเทพ", "กทม", "กระบี่", "กาญจนบุรี", "กาฬสินธุ์", "กำแพงเพชร",
    "ขอนแก่น", "จันทบุรี", "ฉะเชิงเทรา", "ชลบุรี", "ชัยนาท", "ชัยภูมิ", "ชุมพร", "เชียงราย",
    "เชียงใหม่", "ตรัง", "ตราด", "ตาก", "นครนายก", "นครปฐม", "นครพนม", "นครราชสีมา",
    "นครศรีธรรมราช", "นครสวรรค์", "นนทบุรี", "นราธิวาส", "น่าน", "บึงกาฬ", "บุรีรัมย์",
    "ปทุมธานี", "ประจวบคีรีขันธ์", "ปราจีนบุรี", "ปัตตานี", "พระนครศรีอยุธยา", "พะเยา",
    "พังงา", "พัทลุง", "พิจิตร", "พิษณุโลก", "เพชรบุรี", "เพชรบูรณ์", "แพร่", "ภูเก็ต",
    "มหาสารคาม", "มุกดาหาร", "แม่ฮ่องสอน", "ยโสธร", "ยะลา", "ร้อยเอ็ด", "ระนอง", "ระยอง",
    "ราชบุรี", "ลพบุรี", "ลำปาง", "ลำพูน", "เลย", "ศรีสะเกษ", "สกลนคร", "สงขลา", "สตูล",
    "สมุทรปราการ", "สมุทรสงคราม", "สมุทรสาคร", "สระแก้ว", "สระบุรี", "สิงห์บุรี", "สุโขทัย",
    "สุพรรณบุรี", "สุราษฎร์ธานี", "สุรินทร์", "หนองคาย", "หนองบัวลำภู", "อ่างทอง",
    "อำนาจเจริญ", "อุดรธานี", "อุตรดิตถ์", "อุทัยธานี", "อุบลราชธานี"
]

def clean_thai_text(text: str) -> str:
    if not text:
        return ""
    # Normalize unicode am-sara and spacing
    text = text.replace('\u0e33', 'ำ').replace('\u0e32\u0e4d', 'ำ').replace('\u0e4d\u0e32', 'ำ')
    text = text.replace('ตํ', 'ตำ').replace('อํ', 'อำ')
    # Replace thai digits to arabic
    thai_digits = "๐๑๒๓๔๕๖๗๘๙"
    for idx, d in enumerate(thai_digits):
        text = text.replace(d, str(idx))
    # Common OCR name prefix misrecognitions
    text = re.sub(r'(?<![ก-๙])(?:หาาย|บาย|ทาย|มาย|ผาย)\s+(?=[ก-๙])', 'นาย ', text)
    # Common OCR name misrecognitions
    text = text.replace('รุ่งภิพย์', 'รุ่งทิพย์').replace('พืรพัฒน์', 'พีรพัฒน์').replace('เมซี', 'เมฆี')
    text = text.replace('เอลิบคักดิ์', 'เฉลิมศักดิ์').replace('โกรษร', 'ไกรษร')
    text = text.replace('อีรศักดิ์', 'ธีรศักดิ์')
    return text

def parse_tr14_address(text: str) -> Optional[str]:
    """
    Robust TR.14 address extractor handling multi-column OCR artifacts.
    """
    text = clean_thai_text(text)
    
    # Look for explicit Address Block in TR.14
    addr_block = ""
    m_block = re.search(r'(?:ข้อมูลที่อยู่|ที่อยู่ตามทะเบียนบ้าน|ที่อยู่)(.*?)(?:สถานภาพบุคคล|ข้อมูลบิดา|ข้อมูลมารดา|ชื่อ-สกุล|ประวัติ|สิทธิ|วันที[่ขย]้ายเข้า|วันที่ย้ายเข้า|สํานักทะเบียน|สำนักทะเบียน|$)', text, re.DOTALL | re.IGNORECASE)
    if m_block:
        addr_block = m_block.group(1)
    else:
        addr_block = text

    addr_parts = []
    
    # House No
    m_hno = re.search(r'(?:บ้านเลขที่|เลขที่)\s*[:\.]?\s*([0-9]+(?:/[0-9]+)?)', addr_block)
    if m_hno:
        addr_parts.append(f"บ้านเลขที่ {m_hno.group(1)}")
        
    # Moo
    m_moo = re.search(r'(?:หมู่ที่|หมู่|ม\.)\s*[:\.]?\s*([0-9]+)', addr_block)
    if m_moo:
        addr_parts.append(f"หมู่ {m_moo.group(1)}")
        
    # Soi
    m_soi = re.search(r'(?:ซอย|ขอย)\s*[:\.]?\s*([ก-๙a-zA-Z0-9\-_]+)', addr_block)
    if m_soi:
        s_val = m_soi.group(1).strip()
        is_english_noise = bool(re.match(r'^[a-zA-Z]{1,4}$', s_val))
        has_thai_or_num = any('\u0e00' <= c <= '\u0e7f' or c.isdigit() for c in s_val)
        if (
            s_val not in ["ถนน", "ตำบล", "อำเภอ", "จังหวัด", "ไม่มี", "-", "สถานภาพ", "ราษฎร์", "รากฏร์บูรณะ"]
            and not is_english_noise
            and has_thai_or_num
            and len(s_val) >= 2
        ):
            addr_parts.append(f"ซอย {s_val}")
            
    # Road / Thanon
    m_road = re.search(r'(?:ถนน|ถ\.)\s*[:\.]?\s*([ก-๙a-zA-Z0-9\-_]+)', addr_block)
    if m_road:
        r_val = m_road.group(1).strip()
        if r_val not in ["ตำบล", "อำเภอ", "จังหวัด", "ไม่มี", "-", "สถานภาพ"]:
            addr_parts.append(f"ถ.{r_val}")

    # Subdistrict (ตำบล / แขวง / ต.)
    m_sub = re.search(r'(?:ตำบล|ต\.|แขวง)\s*[:\.]?\s*([ก-๙]+)', addr_block)
    if m_sub:
        sub_val = m_sub.group(1).strip()
        sub_val = re.sub(r'^[^\wก-๙]+', '', sub_val)
        if sub_val not in ["อำเภอ", "จังหวัด", "สํานักทะเบียน", "สำนักทะเบียน", "ประเภทบ้าน", "เพศ", "วัน", "เกิด", "ชาย", "หญิง", "ไทย", "สถานภาพ"]:
            addr_parts.append(f"ต.{sub_val}")

    # District (อำเภอ / อ. / เขต)
    m_dist = re.search(r'(?:อำเภอ|อ\.|เขต|อําเกอ)\s*[:\.]?\s*([ก-๙]+)', addr_block)
    if m_dist:
        dist_val = m_dist.group(1).strip()
        dist_val = re.sub(r'^[^\wก-๙]+', '', dist_val)
        if dist_val not in ["จังหวัด", "สํานักทะเบียน", "สำนักทะเบียน", "เพศ", "สถานภาพ", "ไทย"]:
            addr_parts.append(f"อ.{dist_val}")

    # Province (จังหวัด / จ.)
    found_prov = None
    for prov in THAI_PROVINCES:
        if prov in addr_block or prov in text:
            found_prov = prov
            break
            
    if not found_prov:
        m_prov = re.search(r'(?:จังหวัด|จ\.)\s*[:\.]?\s*([ก-๙]+)', addr_block)
        if m_prov:
            p_val = m_prov.group(1).strip()
            if p_val not in ["สถานภาพ", "สถานภาพบุคคล", "สํานักทะเบียน", "สำนักทะเบียน", "เพศ", "ไทย", "วันที่ย้ายเข้า", "วันทีข้ายเข้า"]:
                found_prov = p_val
                
    if found_prov:
        addr_parts.append(f"จ.{found_prov}")

    if addr_parts:
        return " ".join(addr_parts)
    return None

def clean_and_format_thai_address(raw_text: str) -> Optional[str]:
    """
    Standardizes Thai administrative address into 'บ้านเลขที่ X หมู่ Y ต.Z อ.W จ.P'.
    Strictly terminates at Province or postal code, discarding any trailing narrative prose,
    case notes, social media accounts, or other suspects' names.
    """
    if not raw_text:
        return None
    text = clean_thai_text(raw_text)

    # 1. Pre-clean registration form noise and headers
    text = re.sub(r'เลขรหัสประจำบ้าน\s*[\d\-\s]+', '', text)
    text = re.sub(r'สถาน[ะภาพ]+เจ้าบ้าน\s*[^\s]+', '', text)
    text = re.sub(r'สถานภาพบุคคล\s*[^\n\r]+', '', text)
    text = re.sub(r'วันที[่ขย]้ายเข้า[^\n\r]*', '', text)
    text = re.sub(r'สํานักทะเบียน[^\n\r]*', '', text)
    text = re.sub(r'สำนักทะเบียน[^\n\r]*', '', text)
    text = re.sub(r'^(?:ข้อมูลที่อยู่|ที่อยู่ตามทะเบียนบ้าน|ที่อยู่|อยู่บ้านเลขที่|บ้านเลขที่|บ้านเลขที|เลขที่)\s*[:\.]?\s*', '', text).strip()

    # 2. Find the FIRST occurrence of a recognized Thai Province
    first_prov_match = None
    for prov in THAI_PROVINCES:
        m_p = re.search(r'(?:จังหวัด|จ\.)\s*' + re.escape(prov) + r'|\b' + re.escape(prov) + r'\b', text)
        if m_p:
            if first_prov_match is None or m_p.start() < first_prov_match[0].start():
                first_prov_match = (m_p, prov)

    found_prov = None
    if first_prov_match:
        m_p, prov_name = first_prov_match
        found_prov = prov_name
        prov_end = m_p.end()
        # Check if 5-digit postal code immediately follows province
        after_prov = text[prov_end:].strip()
        m_post = re.match(r'^[^\d]*(\d{5})\b', after_prov)
        if m_post:
            prov_end += m_post.end()
        # Strictly truncate at province boundary
        text = text[:prov_end].strip()
    else:
        # Fallback: check if 5-digit postal code exists
        m_post = re.search(r'\b\d{5}\b', text)
        if m_post:
            text = text[:m_post.end()].strip()
        else:
            # Cut off at known case / social keywords if present
            for kw in ['ติดต่อ', 'จับกุม', 'ยาบ้า', 'facebook', 'ไลน์', 'line', 'โทร', 'โทรศัพท์', 'เจ้าหน้าที่', 'ของกลาง']:
                idx_kw = text.find(kw)
                if idx_kw > 0:
                    text = text[:idx_kw].strip()
                    break

    # 3. Extract standard components
    # House No: avoid false positive on long phone numbers or IDs
    m_hno = re.search(r'(?:บ้านเลขที่|บ้านเลขที|เลขที่)?\s*([0-9]{1,4}(?:/[0-9]{1,4})?)', text)
    hno = m_hno.group(1) if m_hno else ""
    if len(hno) > 6 and '/' not in hno:
        hno = ""

    # Moo
    m_moo = re.search(r'(?:หมู่ที่|หมู่ที|หมู่|ม\.)\s*([0-9]+)', text)
    moo = m_moo.group(1) if m_moo else ""

    # Soi
    m_soi = re.search(r'(?:ซอย|ขอย)\s*([ก-๙a-zA-Z0-9\-_]+)', text)
    soi = ""
    if m_soi:
        s_val = m_soi.group(1).strip()
        is_english_noise = bool(re.match(r'^[a-zA-Z]{1,4}$', s_val))
        has_thai_or_num = any('\u0e00' <= c <= '\u0e7f' or c.isdigit() for c in s_val)
        if (
            s_val not in ["ถนน", "ตำบล", "อำเภอ", "จังหวัด", "ไม่มี", "-", "ราษฎร์", "รากฏร์บูรณะ"]
            and not is_english_noise
            and has_thai_or_num
            and len(s_val) >= 2
        ):
            soi = s_val

    # Road
    m_road = re.search(r'(?:ถนน|ถ\.)\s*([ก-๙a-zA-Z0-9\-_]+)', text)
    road = ""
    if m_road:
        r_val = m_road.group(1).strip()
        if r_val not in ["ตำบล", "อำเภอ", "จังหวัด", "ไม่มี", "-", "หัก"]:
            road = r_val

    # Subdistrict
    m_sub = re.search(r'(?:ตำบล|ต\.|แขวง|ดำบล|ต่ําบล)\s*([ก-๙]+)', text)
    sub = ""
    if m_sub:
        sub_val = m_sub.group(1).strip()
        if sub_val not in ["อำเภอ", "จังหวัด", "ไม่มี", "-"]:
            sub = sub_val

    # District
    m_dist = re.search(r'(?:อำเภอ|อ\.|เขต|อําเกอ|อำเกอ|อําเภอ)\s*([ก-๙]+)', text)
    dist = ""
    if m_dist:
        dist_val = m_dist.group(1).strip()
        if dist_val not in ["จังหวัด", "ไม่มี", "-"]:
            dist = dist_val

    parts = []
    if hno:
        parts.append(f"บ้านเลขที่ {hno}")
    if moo:
        parts.append(f"หมู่ {moo}")
    if soi:
        parts.append(f"ซอย {soi}")
    if road:
        parts.append(f"ถ.{road}")
    if sub:
        parts.append(f"ต.{sub}")
    if dist:
        parts.append(f"อ.{dist}")
    if found_prov:
        parts.append(f"จ.{found_prov}")
    elif not found_prov:
        m_pr = re.search(r'(?:จังหวัด|จ\.)\s*([ก-๙]+)', text)
        if m_pr:
            parts.append(f"จ.{m_pr.group(1).strip()}")

    if len(parts) >= 2:
        return " ".join(parts)

    clean_sub = re.sub(r'[\(\),\|]+', ' ', text)
    clean_sub = re.sub(r'\s+', ' ', clean_sub).strip()
    return clean_sub if clean_sub else None

def parse_report_address(text: str) -> Optional[str]:
    """
    Extracts and standardizes address from prose police reports (DOCX / PDF).
    Terminates strictly at Province boundary to prevent narrative prose bleed.
    """
    if not text:
        return None
    cleaned_input = clean_thai_text(text)
    m = re.search(r'(?:อยู่บ้านเลขที่|บ้านเลขที่|บ้านเลขที|ที่อยู่)\s*([0-9/].*?)(?=(?:เลขประจำตัวประชาชน|บัตรประจำตัว|โทรศัพท์|อายุ|\n|$))', cleaned_input)
    if m:
        cand = m.group(1).strip()
        cand = re.sub(r'[\(\),]+$', '', cand).strip()
        return clean_and_format_thai_address(cand)
    # If no explicit prefix matched, search for direct address pattern in text
    return clean_and_format_thai_address(cleaned_input)


def extract_seized_items(text: str) -> Dict[str, Any]:
    """
    Extracts seized narcotics, firearms/ammunition, and cash from police reports or evidence lists.
    Prioritizes official seizure list section over narrative confession text.
    """
    items = {
        "drugs": [],
        "weapons": [],
        "cash": 0.0,
        "total_drug_quantity": 0,
        "summary": [],
        "is_official_seizure": False
    }
    if not text:
        return items

    # 0. Isolate official seizure section from narrative confession / past drug transactions
    m_seizure = re.search(
        r'(?:พร้อมด้วยของกลางดังนี[้]?|พร้อมของกลางดังนี[้]?|รายการของกลางดังนี[้]?|ของกลางดังนี[้]?|ของกลางมีดังนี[้]?|ยึดของกลางดังนี[้]?|ของกลางที่ตรวจยึด(?:ได้)?ดังนี[้]?)'
        r'(.*?)(?:ตำแหน่งที่พบของกลาง|พฤติการณ์ในการจับกุม|พฤติการณ์แห่งการจับกุม|พฤติการณ์|เอกสารประกอบคดี|ก่อนการจับกุม|โดยกล่าวหาว่า|ข้อหา|สิทธิของผู้ถูกจับ|$)',
        text,
        re.DOTALL | re.IGNORECASE
    )
    is_official_seizure = bool(m_seizure)
    items["is_official_seizure"] = is_official_seizure
    seizure_text = m_seizure.group(1) if m_seizure else text

    # 1. Check for explicit grand total summary line first (e.g. รวมของกลางยาบ้าทั้งหมดในคดีนี้ทั้งสิ้น 1,030 เม็ด)
    m_total = re.search(r'(?:รวมของกลาง|ตรวจนับ(?:รวม)?(?:ของกลาง)?|ของกลางรวม)(?:ยาบ้า|ยาเสพติด)?(?:ทั้งหมด)?(?:ในคดีนี้)?(?:ทั้งสิ้น)?\s*[:\.]?\s*([0-9,]+)\s*(เม็ด|กรัม|กก\.)', seizure_text)
    if not m_total:
        m_total = re.search(r'(?:รวมของกลาง|ตรวจนับ(?:รวม)?(?:ของกลาง)?|ของกลางรวม)(?:ยาบ้า|ยาเสพติด)?(?:ทั้งหมด)?(?:ในคดีนี้)?(?:ทั้งสิ้น)?\s*[:\.]?\s*([0-9,]+)\s*(เม็ด|กรัม|กก\.)', text)

    has_explicit_total = False
    explicit_qty = 0
    if m_total:
        try:
            explicit_qty = int(m_total.group(1).replace(',', ''))
            has_explicit_total = True
        except ValueError:
            pass

    # 1. Seized drugs: ยาบ้า, ไอซ์, เฮโรอีน, เคตามีน, กัญชา, ฝิ่น, ยาอี
    # Allow spanning newlines and descriptions up to 350 chars inside official seizure section
    max_span = 350 if is_official_seizure else 80
    drug_pattern = re.finditer(
        r'(ยาบ้า|ยาไอซ์|ไอซ์|เฮโรอีน|เคตามีน|ยาเค|กัญชา|ยาอี|เอ็กซ์ตาซี|ยาเสพติดให้โทษ[^\n\r,]*?)[^\d]{0,' + str(max_span) + r'}?'
        r'(\d{1,3}(?:,\d{3})*(?:\.\d+)?|\d+)\s*'
        r'(เม็ด|กรัม|กิโลกรัม|กก\.|กิโล|ถุง|ห่อ|มัด|หลอด)',
        seizure_text,
        re.IGNORECASE
    )
    summed_drugs = 0
    seen_drug_quantities = set()
    for m in drug_pattern:
        # If not official seizure, skip lines that relate to sting ops, expansion, other persons, or confessions
        if not is_official_seizure:
            start_pos = max(0, m.start() - 60)
            end_pos = min(len(seizure_text), m.end() + 60)
            context_snippet = seizure_text[start_pos:end_pos]
            if any(k in context_snippet for k in ["ขยายผล", "ตรวจยึดเพิ่มเติม", "แยกดำเนินคดี", "ล่อซื้อ", "เคยซื้อ", "ขอซื้อ", "โอนเงิน", "บัญชี", "ขายในราคา", "กระท่อม"]):
                continue

        d_name = m.group(1).strip()
        if 'ยาเสพติด' in d_name and ('ยาบ้า' in d_name or 'เมทแอมเฟตามีน' in d_name or 'ประเภท' in d_name or m.group(3).strip() == 'เม็ด'):
            d_name = 'ยาบ้า'
        qty_str = m.group(2).replace(',', '')
        try:
            qty = float(qty_str) if '.' in qty_str else int(qty_str)
        except ValueError:
            continue
        unit = m.group(3).strip()

        # In fallback mode, do not blindly sum identical quantities of the same drug repeated across paragraphs
        if not is_official_seizure and (d_name, qty, unit) in seen_drug_quantities:
            continue
        seen_drug_quantities.add((d_name, qty, unit))

        items["drugs"].append({
            "name": d_name,
            "quantity": qty,
            "unit": unit
        })
        items["summary"].append(f"{d_name} {m.group(2)} {unit}")
        if d_name in ["ยาบ้า", "ยาอี"] and unit == "เม็ด":
            summed_drugs += int(qty)

    if has_explicit_total and explicit_qty > 0:
        items["total_drug_quantity"] = explicit_qty
        if not any("รวมของกลาง" in s for s in items["summary"]):
            items["summary"].insert(0, f"รวมของกลางยาบ้า {explicit_qty:,} เม็ด")
    else:
        items["total_drug_quantity"] = summed_drugs

    # 2. Seized weapons & ammunition: ปืน, อาวุธปืน, ลูกซอง, กระสุน
    # If official seizure section exists, strictly only look within official seizure section
    weapon_text = seizure_text if is_official_seizure else (seizure_text if any(k in seizure_text for k in ["ปืน", "อาวุธ", "กระสุน"]) else text)
    
    # Firearms (unit กระบอก)
    gun_pattern = re.finditer(
        r'(อาวุธปืน(?:[^\d\n\r,]{0,40}(?:ไทยประดิษฐ์|สงคราม|สั้น|ยาว|ลูกซอง|พก|ออโตเมติก|รีวอลเวอร์))?|อาวุธปืน|'
        r'ปืน(?:[^\d\n\r,]{0,40}(?:ไทยประดิษฐ์|สงคราม|สั้น|ยาว|ลูกซอง|พก|ออโตเมติก|รีวอลเวอร์))?|ปืน)'
        r'(?:[^\d\n\r,]{0,40}?(?:ขนาด\s*(?:เบอร์\s*)?[\d\.]+\s*(?:มม\.|นิ้ว|มิลลิเมตร|เกจ)?))?'
        r'(?:[^\d\n\r,]{0,20}?จำนวน)?\s*'
        r'(\d{1,4})\s*'
        r'(กระบอก)',
        weapon_text
    )
    # Ammunition (unit นัด, ลูก)
    ammo_pattern = re.finditer(
        r'(เครื่องกระสุนปืนลูกซอง|เครื่องกระสุนปืน|เครื่องกระสุน|ลูกกระสุนปืน|ลูกกระสุน|กระสุนปืน|กระสุน)'
        r'(?:[^\d\n\r,]{0,40}?(?:ขนาด\s*(?:เบอร์\s*)?[\d\.]+\s*(?:มม\.|นิ้ว|มิลลิเมตร|เกจ)?))?'
        r'(?:[^\d\n\r,]{0,20}?จำนวน)?\s*'
        r'(\d{1,4})\s*'
        r'(นัด|ลูก)',
        weapon_text
    )

    seen_weapons = set()
    for m in list(gun_pattern) + list(ammo_pattern):
        if not is_official_seizure:
            start_pos = max(0, m.start() - 60)
            end_pos = min(len(weapon_text), m.end() + 60)
            context_snippet = weapon_text[start_pos:end_pos]
            if any(k in context_snippet for k in ["ขยายผล", "ตรวจยึดเพิ่มเติม", "แยกดำเนินคดี", "ล่อซื้อ", "กระท่อม"]):
                continue
        w_name = re.sub(r'\s+', ' ', m.group(1).strip())
        qty_str = m.group(2).replace(',', '')
        try:
            qty = int(qty_str)
        except ValueError:
            continue
        unit = m.group(3).strip()

        weapon_key = (w_name, qty, unit)
        if weapon_key in seen_weapons:
            continue
        seen_weapons.add(weapon_key)

        items["weapons"].append({
            "name": w_name,
            "quantity": qty,
            "unit": unit
        })
        items["summary"].append(f"{w_name} {qty} {unit}")

    # 3. Seized cash: เงินสด
    cash_text = seizure_text if is_official_seizure else (seizure_text if "เงินสด" in seizure_text else text)
    cash_pattern = re.search(
        r'(?:เงินสด|ธนบัตร)[^\d\n\r]{0,80}?'
        r'(\d{1,3}(?:,\d{3})*(?:\.\d{2})?|\d+)\s*'
        r'(?:บาท)?',
        cash_text
    )
    if cash_pattern:
        if not is_official_seizure:
            start_pos = max(0, cash_pattern.start() - 60)
            end_pos = min(len(cash_text), cash_pattern.end() + 60)
            context_snippet = cash_text[start_pos:end_pos]
            if not any(k in context_snippet for k in ["ขยายผล", "ล่อซื้อ", "โอนเงิน", "ซื้อยาบ้า", "ขายยาบ้า"]):
                try:
                    cash_val = float(cash_pattern.group(1).replace(',', ''))
                    if cash_val > 0:
                        items["cash"] = cash_val
                        items["summary"].append(f"เงินสด {cash_val:,.0f} บาท")
                except ValueError:
                    pass
        else:
            try:
                cash_val = float(cash_pattern.group(1).replace(',', ''))
                if cash_val > 0:
                    items["cash"] = cash_val
                    items["summary"].append(f"เงินสด {cash_val:,.0f} บาท")
            except ValueError:
                pass

    # Deduplicate items["summary"] preserving order, omitting sub-packaging units if total pills is present
    clean_summary = []
    seen_sum = set()
    drugs_with_pills = {d["name"] for d in items["drugs"] if d["unit"] == "เม็ด"}
    for s in items["summary"]:
        skip_pkg = False
        for dname in drugs_with_pills:
            if s.startswith(dname) and any(pkg in s for pkg in ["ถุง", "มัด", "ห่อ", "หลอด"]):
                skip_pkg = True
                break
        if skip_pkg:
            continue
        if s not in seen_sum:
            seen_sum.add(s)
            clean_summary.append(s)
    items["summary"] = clean_summary

    return items


def extract_vehicle_plates(text: str) -> List[str]:
    """
    Extracts Thai vehicle license plates e.g. 1กก 9999 บุรีรัมย์ or กข 1234 กทม.
    """
    plates = []
    if not text:
        return plates
    plate_pattern = re.finditer(
        r'(?:ทะเบียน(?:รถ)?|ป้ายทะเบียน|หมายเลขทะเบียน)?\s*'
        r'(?<![ก-๙a-zA-Z0-9])([0-9ก-ฮ]{1,2}\s*[ก-ฮ]{1,2})\s*[-–\s]?\s*(\d{1,4})\b\s*'
        r'(?:จังหวัด|จ\.)?\s*([ก-๙]{2,})?',
        text
    )
    NON_PLATE_PREFIXES = {'อง', 'อัน', 'เครื่อง', 'คัน', 'ชิ้น', 'ใบ', 'กระบอก', 'นัด', 'เม็ด', 'ห่อ', 'ถุง', 'มัด'}
    seen = set()
    for m in plate_pattern:
        pfx = m.group(1).replace(' ', '')
        num = m.group(2).strip()
        prov = m.group(3).strip() if m.group(3) else ""
        if not re.search(r'[ก-ฮ]', pfx):
            continue
        if pfx in NON_PLATE_PREFIXES:
            continue
        if len(num) == 4 and int(num) in range(2500, 2600):
            continue
        full_plate = f"{pfx} {num} {prov}".strip()
        if full_plate not in seen:
            seen.add(full_plate)
            plates.append(full_plate)
    return plates


def extract_vehicle_registration_doc(file_bytes: bytes, text: Optional[str] = None) -> Dict[str, Any]:
    """
    Extracts structured vehicle data from DLT Vehicle Registration Documents / Green Book copies.
    Handles corporate owner exclusion (banks, leasing companies, corporations).
    """
    res: Dict[str, Any] = {
        "is_vehicle_doc": False,
        "plate": None,
        "province": None,
        "vehicle_type": None,
        "brand": None,
        "model": None,
        "color": None,
        "chassis_no": None,
        "engine_no": None,
        "owner_name": None,
        "owner_id": None,
        "finance_company": None,
        "is_corporate_owner": False,
        "raw_text": ""
    }
    try:
        if not text:
            raw_img = Image.open(io.BytesIO(file_bytes))
            clean_img = _remove_red_markup(raw_img)
            text = pytesseract.image_to_string(clean_img, lang="tha+eng")
        
        if not text:
            return res

        cleaned = clean_thai_text(text)
        res["raw_text"] = cleaned

        # Detect DLT indicators
        is_dlt = any(k in cleaned for k in [
            "รายการจดทะเบียน", "คู่มือจดทะเบียน", "ใบคู่มือจดทะเบียน", "กรมการขนส่งทางบก",
            "เลขทะเบียน", "ประเภท", "ยี่ห้อรถ", "เลขตัวรถ", "เลขเครื่องยนต์", "ผู้ถือกรรมสิทธิ์", "ผู้ครอบครอง"
        ])
        
        # 1. Plate & Province
        m_plate = re.search(r'(?:เลขทะเบียน|ทะเบียนรถ|ทะเบียน|หมายเลขทะเบียน)\s*[:\.]?\s*([1-9]?[ก-ฮ]{1,2}\s*\d{1,4})', cleaned)
        if not m_plate:
            m_plate = re.search(r'\b([1-9]?[ก-ฮ]{1,2}\s*\d{1,4})\b', cleaned)
        
        if m_plate:
            p_val = m_plate.group(1).strip()
            if re.search(r'[ก-ฮ]', p_val):
                res["plate"] = p_val
                if not is_dlt and any(k in cleaned for k in ["รถเก๋ง", "รถยนต์", "รถกระบะ", "ยี่ห้อ", "สี"]):
                    is_dlt = True

        # Province
        for prov in THAI_PROVINCES:
            if prov in cleaned:
                res["province"] = prov
                break
        if not res["province"]:
            for abbr, full_name in [("กทม", "กรุงเทพมหานคร"), ("บร", "บุรีรัมย์"), ("นม", "นครราชสีมา")]:
                if re.search(r'\b' + abbr + r'\b', cleaned):
                    res["province"] = full_name
                    break

        # 2. Vehicle Type (ประเภท)
        m_type = re.search(r'(?:ประเภท|ชนิด)\s*[:\.]?\s*([^\n\r]+?)(?=(?:ยี่ห้อ|แบบ|สี|เลข|$))', cleaned)
        if m_type:
            res["vehicle_type"] = m_type.group(1).strip()
        else:
            for v_type in ["รถยนต์นั่งส่วนบุคคลไม่เกิน 7 คน", "รถยนต์บรรทุกส่วนบุคคล", "รถจักรยานยนต์", "รถเก๋ง", "รถกระบะ", "รถตู้", "รถบรรทุก"]:
                if v_type in cleaned:
                    res["vehicle_type"] = v_type
                    break

        # 3. Brand & Model (ยี่ห้อ / รุ่น)
        m_brand = re.search(r'(?:ยี่ห้อรถ|ยี่ห้อ)\s*[:\.]?\s*([a-zA-Zก-๙]+)', cleaned)
        if m_brand:
            res["brand"] = m_brand.group(1).strip().upper()
        else:
            for b in ["TOYOTA", "HONDA", "ISUZU", "MITSUBISHI", "NISSAN", "MAZDA", "FORD", "MG", "SUZUKI", "YAMAHA", "BMW", "BENZ"]:
                if b in cleaned.upper():
                    res["brand"] = b
                    break

        m_model = re.search(r'(?:แบบ|รุ่น)\s*[:\.]?\s*([a-zA-Z0-9\-_]+)', cleaned)
        if m_model:
            res["model"] = m_model.group(1).strip().upper()

        # 4. Color (สี)
        m_color = re.search(r'(?:สีรถ|สี)\s*[:\.]?\s*([ก-๙]+)', cleaned)
        if m_color and len(m_color.group(1).strip()) >= 2:
            res["color"] = m_color.group(1).strip()
        else:
            for c in ["สีขาว", "สีดำ", "สีเทา", "สีบรอนซ์", "สีเงิน", "สีแดง", "สีน้ำเงิน", "สีเขียว", "สีเหลือง", "สีส้ม", "สีน้ำตาล"]:
                if c in cleaned:
                    res["color"] = c.replace("สี", "")
                    break

        # 5. Chassis & Engine Number
        m_chassis = re.search(r'(?:เลขตัวรถ|เลขคัสซี|เลขตัวถัง)\s*[:\.]?\s*([a-zA-Z0-9]+)', cleaned)
        if m_chassis:
            res["chassis_no"] = m_chassis.group(1).strip()

        m_engine = re.search(r'(?:เลขเครื่องยนต์|เลขเครื่อง)\s*[:\.]?\s*([a-zA-Z0-9]+)', cleaned)
        if m_engine:
            res["engine_no"] = m_engine.group(1).strip()

        # 6. Owner & Corporate Filtering
        PREFIXES = r'(?:(?<![ก-๙a-zA-Z])(?:นาย|หาาย|บาย|ทาย|มาย|ผาย|นางสาว|มางสาว|บา่งสาว|น\.ส\.|นาง(?!รอง)|เด็กชาย|เด็กหญิง|ด\.ช\.|ด\.ญ\.))'
        m_owner = re.search(r'(?:ผู้ถือกรรมสิทธิ์|ผู้ครอบครอง|ชื่อผู้ถือกรรมสิทธิ์|เจ้าของ)\s*[:\.]?\s*([^\n\r]+?)(?=(?:เลขประจำตัว|วันเดือนปี|ที่อยู่|\n|$))', cleaned)
        
        raw_owner_name = None
        if m_owner:
            raw_owner_name = m_owner.group(1).strip()
        else:
            m_p_owner = re.search(r'(' + PREFIXES + r'\s*[ก-๙]{2,}\s+[ก-๙]{2,})', cleaned)
            if m_p_owner:
                raw_owner_name = m_p_owner.group(1).strip()

        CORPORATE_REGEX = re.compile(
            r'(?:บริษัท|บจก\.|บมจ\.|ห้างหุ้นส่วน|หจก\.|ลิสซิ่ง|แคปปิตอล|ธนาคาร|ไฟแนนซ์|leasing|bank|capital|ลีสซิ่ง|โตโยต้า\s*ลีสซิ่ง|ฮอนด้า\s*ลีสซิ่ง)',
            re.IGNORECASE
        )

        if raw_owner_name:
            if CORPORATE_REGEX.search(raw_owner_name):
                res["is_corporate_owner"] = True
                res["finance_company"] = raw_owner_name
                res["owner_name"] = None
            else:
                res["owner_name"] = normalize_thai_name_candidate(raw_owner_name)

        # Owner Citizen ID
        m_cid = re.search(r'(?:เลขประจำตัวประชาชน|เลขประจำตัวผู้เสียภาษีอากร|เลขบัตรประชาชน)\s*[:\.]?\s*([0-9\s\-]{13,17})', cleaned)
        if m_cid:
            clean_id = re.sub(r'\D', '', m_cid.group(1))
            if len(clean_id) == 13 and not res["is_corporate_owner"]:
                res["owner_id"] = clean_id

        res["is_vehicle_doc"] = is_dlt and (bool(res["plate"]) or bool(res["chassis_no"]))
        return res
    except Exception as e:
        logger.error(f"Failed to extract vehicle registration document: {e}")
        return res


DISALLOWED_SOCIAL_WORDS = {
    'ชื่อ', 'ผู้ใช้', 'โปรไฟล์', 'แอด', 'กลุ่ม', 'official', 'com', 'www', 'http', 'https',
    'ไม่ทราบ', 'แอพพลิเคชั่น', 'แอพพลิเคชัน', 'แอปพลิเคชัน', 'ดังกล่าวจริง', 'ไลน์', 'เฟส', 'เฟสบุ๊ค',
    'ดังกล่าว', 'จริง', 'โทรศัพท์', 'โทรศัพท์มือถือ', 'และ', 'หรือ', 'กับ', 'แอพ', 'แอป'
}

NARRATIVE_JUNK_WORDS = [
    'เจ้าหน้าที่', 'ตำรวจ', 'ดังกล่าว', 'ปรากฏ', 'ข้อความ', 'ตรวจสอบ', 'จับกุม', 'ยอมรับ',
    'โทรศัพท์', 'รวบรวม', 'หลักฐาน', 'โทรศัพท์มือถือ', 'ของกลาง', 'แอพพลิเคชั่น', 'แอพพลิเคชัน', 'แอปพลิเคชัน'
]

NARRATIVE_STOP_PATTERN = re.compile(
    r'\s+(?:'
    r'ใช้|ติดต่อ|โทร|และ|โดย|เพื่อ|ส่ง|ซึ่ง|เมื่อ|พี่|นาย|นาง|น\.ส\.|เป็น|จึง|จริง|'
    r'ระหว่าง|ปรากฏ|สอบถาม|ขอให้|ตกลง|ทาง|ไลน์|เฟส|เฟซ|facebook|line|tiktok'
    r')',
    re.IGNORECASE
)


def clean_social_account(acc: str) -> Optional[str]:
    if not acc:
        return None
    acc = acc.strip(" \t\n\r\"'”’“‘«»()[]{}:;,")
    m_stop = NARRATIVE_STOP_PATTERN.search(acc)
    if m_stop:
        acc = acc[:m_stop.start()].strip()
    acc = re.sub(r'^(?:ของนาย|ของนางสาว|ของนาง|ของ)\s*', '', acc).strip()
    acc = acc.strip(" \t\n\r\"'”’“‘«»()[]{}:;,")
    if not acc or len(acc) < 2 or len(acc) > 40:
        return None
    if acc in DISALLOWED_SOCIAL_WORDS:
        return None
    if any(j in acc for j in NARRATIVE_JUNK_WORDS):
        return None
    if acc.startswith(('และ ', 'หรือ ', 'กับ ', 'ของ ', 'ที่ ', 'ได้ ')):
        return None
    return acc


def extract_social_media(text: str) -> List[Dict[str, str]]:
    """
    Extracts social media accounts (LINE, Facebook, TikTok).
    """
    results = []
    if not text:
        return results
    seen = set()

    # 1. Quoted platform account names e.g. Facebook “เสี่ย แบงค์”, LINE “ปี๊บ อะ จิง โลยๆๆ”
    quoted_pattern = re.finditer(
        r'(?:แอพพลิเคชั่น|แอพพลิเคชัน|แอปพลิเคชัน)?\s*'
        r'(?:(LINE|ไลน์)|(Facebook|เฟซบุ๊ก|เฟสบุ๊ก|เฟซบุ๊ค|เฟสบุ๊ค|เฟซ|เฟส|FB)|(TikTok|ติ๊กต๊อก))\s*'
        r'(?:ID|ไอดี|ชื่อบัญชี|บัญชีชื่อ|ชื่อโปรไฟล์|โปรไฟล์|ชื่อ)?\s*[:\s]*'
        r'(?:[“"]([^”"\r\n]{1,60})[”"]|[‘\']([^’\'\r\n]{1,60})[’\'])',
        text,
        re.IGNORECASE
    )
    for m in quoted_pattern:
        raw_plat = m.group(1) or m.group(2) or m.group(3)
        plat = 'LINE' if re.search(r'line|ไลน์', raw_plat, re.I) else ('Facebook' if re.search(r'face|เฟส|เฟซ', raw_plat, re.I) else 'TikTok')
        raw_val = m.group(4) or m.group(5)
        acc = clean_social_account(raw_val)
        if acc and (plat, acc) not in seen:
            seen.add((plat, acc))
            results.append({"platform": plat, "account": acc})

    # 2. Contextual quoted accounts e.g. บัญชีชื่อ “รอด”
    ctx_quoted = re.finditer(
        r'(?:บัญชีชื่อ|ชื่อบัญชี|บัญชีผู้ใช้|บัญชี)\s*[:\s]*'
        r'(?:[“"]([^”"\r\n]{1,60})[”"]|[‘\']([^’\'\r\n]{1,60})[’\'])',
        text,
        re.IGNORECASE
    )
    for m in ctx_quoted:
        raw_val = m.group(1) or m.group(2)
        acc = clean_social_account(raw_val)
        if acc and not any(acc == r["account"] for r in results):
            start_idx = max(0, m.start() - 120)
            surrounding = text[start_idx:m.end() + 60].lower()
            plat = 'Facebook' if any(k in surrounding for k in ['facebook', 'เฟส', 'เฟซ']) else ('LINE' if any(k in surrounding for k in ['line', 'ไลน์']) else 'Social')
            seen.add((plat, acc))
            results.append({"platform": plat, "account": acc})

    # 3. Unquoted named accounts e.g. LINE ชื่อ มังกร, Facebook ชื่อ Fino wittaya
    unquoted_pattern = re.finditer(
        r'(?:(LINE|ไลน์)|(Facebook|เฟซบุ๊ก|เฟสบุ๊ก|เฟซบุ๊ค|เฟสบุ๊ค|เฟซ|เฟส|FB)|(TikTok|ติ๊กต๊อก))\s*'
        r'(?:ID|ไอดี|ชื่อบัญชี|บัญชีชื่อ|โปรไฟล์)?\s*ชื่อ\s*[:\s]*'
        r'([^\s“"‘\'«\r\n,](?:(?!\s+(?:ใช้|ติดต่อ|ทาง|โทร|และ|เพื่อ|ของ|โดย|ส่ง|พี่|นาย|นาง|ตกลง|ขอให้|ไลน์|เฟส|facebook|line|tiktok))[^\r\n,]){0,35})',
        text,
        re.IGNORECASE
    )
    for m in unquoted_pattern:
        raw_plat = m.group(1) or m.group(2) or m.group(3)
        plat = 'LINE' if re.search(r'line|ไลน์', raw_plat, re.I) else ('Facebook' if re.search(r'face|เฟส|เฟซ', raw_plat, re.I) else 'TikTok')
        acc = clean_social_account(m.group(4))
        if acc and (plat, acc) not in seen and not any(acc == r["account"] for r in results):
            seen.add((plat, acc))
            results.append({"platform": plat, "account": acc})

    # 4. Explicit LINE ID / alphanumeric handle
    line_id_pattern = re.finditer(
        r'(?:LINE\s*ID|Line\s*ID|ไลน์ไอดี|ไอดีไลน์|ID\s*LINE)\s*[:\s]\s*([a-zA-Z0-9._\-@]{3,30})',
        text,
        re.IGNORECASE
    )
    for m in line_id_pattern:
        acc = m.group(1).strip()
        if acc not in seen and len(acc) >= 3 and acc not in DISALLOWED_SOCIAL_WORDS:
            seen.add(('LINE', acc))
            results.append({"platform": "LINE", "account": acc})

    return results


def extract_case_info(text: str) -> Dict[str, Any]:
    """
    Extracts criminal case reference, arrest date, and incident location.
    """
    info = {
        "case_number": None,
        "arrest_date": None,
        "incident_location": None,
        "district": None,
        "subdistrict": None
    }
    if not text:
        return info
    m_case = re.search(r'(?:คดีอาญาที่|คดีที่|สำนวนคดีที่)\s*([0-9]+/[0-9]{4})', text)
    if m_case:
        info["case_number"] = m_case.group(1).strip()
    m_date = re.search(r'(?:เมื่อวันที่|วันที่จับกุม|วันที่)\s*(\d{1,2}\s*(?:ม\.ค\.|ก\.พ\.|มี\.ค\.|เม\.ย\.|พ\.ค\.|มิ\.ย\.|ก\.ค\.|ส\.ค\.|ก\.ย\.|ต\.ค\.|พ\.ย\.|ธ\.ค\.|มกราคม|กุมภาพันธ์|มีนาคม|เมษายน|พฤษภาคม|มิถุนายน|กรกฎาคม|สิงหาคม|กันยายน|ตุลาคม|พฤศจิกายน|ธันวาคม)\s*(?:พ\.ศ\.)?\s*\d{4})', text)
    if m_date:
        info["arrest_date"] = m_date.group(1).strip()

    # Extract actual scene location (excluding arresting unit address context)
    m_scene = re.search(r'(?:สถานที่(?:เกิดเหตุ|จับกุม)|จับกุมได้ที่|ตรวจค้นพบบริเวณ)\s*([^\n\r]{5,100})', text)
    if m_scene:
        scene_raw = m_scene.group(1).strip()
        info["incident_location"] = scene_raw
        m_dist = re.search(r'(?:อ\.|อำเภอ)\s*([^\s,]+)', scene_raw)
        if m_dist:
            info["district"] = m_dist.group(1).strip()
        m_sub = re.search(r'(?:ต\.|ตำบล)\s*([^\s,]+)', scene_raw)
        if m_sub:
            info["subdistrict"] = m_sub.group(1).strip()
    return info


def _get_gemini_api_key() -> Optional[str]:
    """
    Retrieves Gemini API key from environment variables or secure local secret files.
    """
    key = os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")
    if key and key.strip():
        return key.strip()

    secret_candidates = [
        Path("/app/config/gemini_api_key.secret"),
        Path("/data/gemini_api_key.secret"),
        Path(__file__).parent.parent / "config" / "gemini_api_key.secret",
        Path("d:/BPP-VMS/frigate-nvr/nas-bot-service/config/gemini_api_key.secret")
    ]
    for sp in secret_candidates:
        try:
            if sp.exists():
                val = sp.read_text(encoding="utf-8").strip()
                if val:
                    return val
        except Exception:
            pass
    return None


def extract_with_gemini_vision(image_bytes: bytes) -> Optional[Dict[str, Any]]:
    """
    Uses Google Gemini Vision API to accurately extract Thai TR.14 civil registration and arrest records.
    Returns structured dict with suspect_name, id_card, address, phone_number, age, role_tag.
    Gracefully returns None if API key is not configured or request fails.
    """
    api_key = _get_gemini_api_key()
    if not api_key:
        return None

    model_names = [os.getenv("GEMINI_VISION_MODEL", "gemini-flash-latest"), "gemini-flash-lite-latest"]
    b64_img = base64.b64encode(image_bytes).decode("utf-8")

    prompt = (
        "วิเคราะห์ภาพเอกสารราชการไทย (เช่น ทร.14, ทะเบียนบ้าน, บัตรประชาชน, บันทึกจับกุม หรือหมายจับ) "
        "สกัดข้อมูลบุคคลเป้าหมาย/ผู้ต้องหา ออกมาในรูปแบบ JSON ตาม schema นี้เท่านั้น:\n"
        "{\n"
        '  "is_identity_document": true, // เป็นเอกสารราชการ/บัตรประชาชน/ทะเบียนบ้าน/หมายจับ ที่ระบุชื่อบุคคลหรือไม่ หากเป็นภาพถ่ายเจ้าหน้าที่ตำรวจ/ภาพของกลาง/ภาพที่เกิดเหตุ ให้ตอบ false\n'
        '  "suspect_name": "ชื่อ-นามสกุลจริงพร้อมคำนำหน้า สกัดเฉพาะชื่อจริงและนามสกุลจริง ห้ามนำฉายาหรือชื่อเล่น (เช่นคำว่า หรือ...) มาใส่เด็ดขาด ห้ามมีคำว่า ข้อหา, ผู้ต้องหา หรือวลีอย่าง \'ทั้งสองราย\', \'พร้อมของกลาง\' ปะปนมาในชื่อ ตรวจสอบตัวสะกด เช่น ฆ กับ ซ",\n'
        '  "id_card": "เลขประจำตัวประชาชน 13 หลัก (เฉพาะตัวเลข)",\n'
        '  "address": "ที่อยู่ตามทะเบียนบ้านแบบเต็ม เช่น หมู่ 1 ต.สะเตา อ.นางรอง จ.บุรีรัมย์",\n'
        '  "age": "อายุ (ตัวเลขเท่านั้น หรือ null)",\n'
        '  "phone_number": "เบอร์โทรศัพท์ (10 หลัก ถ้ามี หรือ null)",\n'
        '  "role_tag": "บทบาท เช่น ผู้ต้องหา หรือ บุคคลเป้าหมาย"\n'
        "}\n"
        "คำเตือนเด็ดขาด:\n"
        "1. ห้ามนำชื่อหน่วยงานหรือป้ายตั้งโต๊ะของเจ้าหน้าที่ตำรวจ เช่น 'ศป.ปส.อ.นางรอง', 'ชปส.ร้อย ตชด.215', 'ร้อย.อส.', 'กอ.รมน.', 'ฉก.ทพ.26', 'ชป.ปส.ภ.จว.บร.' มาใส่เป็นชื่อบุคคลเด็ดขาด\n"
        "2. หากในภาพไม่มีเอกสารหรือป้ายชื่อระบุตัวผู้ต้องหา มีแต่เจ้าหน้าที่ตำรวจและของกลาง ให้ตอบ suspect_name เป็น null และ id_card เป็น null และ is_identity_document เป็น false\n"
        "3. หากเป็นเอกสาร ทร.14 ให้ดูที่ชื่อตัวและชื่อสกุลของเจ้าของรายการเป็นหลัก ตรวจสอบตัวสะกดภาษาไทยให้ถูกต้อง เช่น ฆ กับ ซ\n"
        "ตอบกลับเป็น JSON เท่านั้น ไม่มี markdown syntax"
    )

    payload = {
        "contents": [
            {
                "role": "user",
                "parts": [
                    {"text": prompt},
                    {
                        "inlineData": {
                            "mimeType": "image/jpeg",
                            "data": b64_img
                        }
                    }
                ]
            }
        ],
        "generationConfig": {
            "temperature": 0.1,
            "responseMimeType": "application/json"
        }
    }

    for model_name in model_names:
        url = f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent"
        try:
            req = urllib.request.Request(
                url,
                data=json.dumps(payload).encode("utf-8"),
                headers={"Content-Type": "application/json", "x-goog-api-key": api_key},
                method="POST"
            )
            with urllib.request.urlopen(req, timeout=15) as resp:
                data = json.loads(resp.read().decode("utf-8"))
                candidates = data.get("candidates", [])
                if candidates and "content" in candidates[0]:
                    parts = candidates[0]["content"].get("parts", [])
                    if parts and "text" in parts[0]:
                        raw_text = parts[0]["text"].strip()
                        if raw_text.startswith("```json"):
                            raw_text = raw_text[7:]
                        if raw_text.startswith("```"):
                            raw_text = raw_text[3:]
                        if raw_text.endswith("```"):
                            raw_text = raw_text[:-3]
                        res_json = json.loads(raw_text.strip())
                        s_name = res_json.get("suspect_name")
                        if s_name:
                            s_name = clean_thai_text(s_name).replace("เมซี", "เมฆี").strip()
                            s_name = normalize_thai_name_candidate(s_name)
                        cid = res_json.get("id_card")
                        if cid and (len(cid) != 13 or cid.startswith('9') or cid.startswith('000')):
                            cid = None

                        if s_name or cid:
                            logger.info(f"Gemini Vision ({model_name}) successfully extracted: {s_name} ({cid})")
                            return {
                                "suspect_name": s_name,
                                "id_card": cid,
                                "address": res_json.get("address"),
                                "age": str(res_json.get("age")) if res_json.get("age") else None,
                                "phone_number": res_json.get("phone_number"),
                                "all_phones": [res_json["phone_number"]] if res_json.get("phone_number") else [],
                                "role_tag": res_json.get("role_tag") or "ผู้ต้องหา",
                                "area_tag": None,
                                "drug_quantity": 0,
                                "extracted_by": f"gemini_vision_{model_name}"
                            }
        except Exception as e:
            logger.warning(f"Gemini Vision extraction failed with {model_name}: {e}")
            continue

    return None


def extract_case_with_gemini_llm(
    text: Optional[str] = None,
    image_bytes: Optional[bytes] = None,
    image_list: Optional[List[bytes]] = None,
    filename: str = ""
) -> Optional[Dict[str, Any]]:
    """
    Uses Google Gemini Flash to accurately extract structured case info, multi-suspect rosters,
    seized contraband items, and scene locations from arrest records or document scans.
    Acts as a Smart Fallback & Full Quality Check engine.
    """
    api_key = _get_gemini_api_key()
    if not api_key:
        return None

    model_names = ["gemini-flash-lite-latest", os.getenv("GEMINI_CASE_MODEL", "gemini-flash-latest")]

    prompt = (
        "คุณคือระบบ AI ผู้เชี่ยวชาญการวิเคราะห์และสกัดข้อมูลจากเอกสารบันทึกการจับกุม/รายงานประจำวัน/หมายจับของเจ้าหน้าที่ตำรวจและตำรวจตระเวนชายแดน (ตชด.)\n"
        "จงอ่านข้อความหรือภาพเอกสารที่แนบมานี้อย่างละเอียด แล้วสกัดข้อมูลที่เกี่ยวข้องออกมาเป็น JSON ตาม schema ดังต่อไปนี้เท่านั้น:\n"
        "{\n"
        '  "suspects": [\n'
        "    {\n"
        '      "name": "ชื่อ-นามสกุลบุคคลจริงๆ พร้อมคำนำหน้า ห้ามนำวลีแปลกปลอมเช่น \'ทั้งสองราย พร้อมตรวจยึดยาบ้า\', \'พร้อมของกลาง\' มาใส่เป็นชื่อเด็ดขาด และห้ามมีคำว่า ผู้ต้องหา, ผู้เสพ หรือข้อหาปะปนในชื่อ",\n'
        '      "id_card": "เลขประจำตัวประชาชน 13 หลัก (เฉพาะตัวเลข หากไม่พบให้ใส่ null)",\n'
        '      "age": "อายุ (ตัวเลขเท่านั้น หรือ null)",\n'
        '      "address": "ที่อยู่ตามทะเบียนบ้าน/ที่พักอาศัย หรือ null",\n'
        '      "role": "บทบาท เช่น ผู้ต้องหาที่ 1, ผู้ต้องหา, หรือ บุคคลเป้าหมาย"\n'
        "    }\n"
        "  ],\n"
        '  "seized_items": {\n'
        '    "meth_pills": 0, // จำนวนยาบ้า (เม็ด) ตัวเลขจำนวนเต็ม\n'
        '    "ice_grams": 0.0, // จำนวนไอซ์ (กรัม) ตัวเลขทศนิยม\n'
        '    "heroin_grams": 0.0, // เฮโรอีน (กรัม)\n'
        '    "firearms": 0, // อาวุธปืน (กระบอก)\n'
        '    "ammunition": 0, // เครื่องกระสุน (นัด)\n'
        '    "cash_thb": 0.0, // เงินสด (บาท)\n'
        '    "vehicles": [] // ป้ายทะเบียนและยานพาหนะ เช่น ["กข 9999 บุรีรัมย์"] ห้ามนำข้อความที่เป็น วัน/เดือน/ปี มาใส่เป็นป้ายทะเบียนรถอย่างเด็ดขาด\n'
        "  },\n"
        '  "crime_scene": {\n'
        '    "location": "สถานที่เกิดเหตุจริง เช่น บริเวณริมถนนสายโนนสุวรรณ-หนองกี่ ต.โนนสุวรรณ อ.โนนสุวรรณ จ.บุรีรัมย์",\n'
        '    "district": "อำเภอที่เกิดเหตุจริง (ห้ามใส่ที่ตั้งหน่วยจับกุม เช่น ร้อย ตชด.215 อ.นางรอง เด็ดขาด)",\n'
        '    "subdistrict": "ตำบลที่เกิดเหตุจริง",\n'
        '    "province": "จังหวัดที่เกิดเหตุ"\n'
        "  },\n"
        '  "arrest_unit": "หน่วยงานที่ทำการจับกุม เช่น ร้อย ตชด.215",\n'
        '  "case_number": "เลขคดีอาญาหรือเลขบันทึกจับกุม (ถ้ามี หรือ null)",\n'
        '  "arrest_date": "วันเวลาที่เกิดเหตุ/จับกุม หรือ null",\n'
        '  "behavior": "สรุปพฤติการณ์สั้นๆ 1-2 ประโยค"\n'
        "}\n"
        "คำเตือนสำคัญ: แยกแยะระหว่าง 'สถานที่เกิดเหตุ' กับ 'ที่ตั้งหน่วยจับกุม' ให้เด็ดขาด อย่าระบุ อ.นางรอง หากเหตุเกิดที่อำเภออื่น ตอบกลับเป็น JSON ล้วนเท่านั้น ไม่มี markdown syntax\n"
        "คำเตือนพิเศษเรื่องป้ายทะเบียน: ห้ามอ่านป้ายทะเบียนจากภาพถ่ายยานพาหนะโดยตรงเด็ดขาด! ให้อ่านเฉพาะข้อความที่ปรากฏใน 'เอกสาร/กระดาษ' เท่านั้น หากภาพที่แนบมาไม่ใช่รูปถ่ายเอกสารหรือกระดาษ ให้ตอบ vehicles เป็น [] เสมอ เพื่อป้องกันความผิดพลาด\n"
        "คำเตือนพิเศษเรื่องของกลาง (seized_items): ให้ระบุเฉพาะของกลางที่ตรวจยึดได้จากตัวหรือที่พักของผู้ต้องหารายนี้โดยตรง (โดยเฉพาะในส่วน 'พร้อมด้วยของกลางดังนี้' หรือป้ายตรวจยึดของกลาง) ห้ามนำของกลางจากการสืบสวนขยายผลในจุดอื่น (เช่น ตรวจยึดยาบ้าเพิ่มเติม 5,000 เม็ดที่กระท่อม) หรือของกลางของบุคคลอื่นที่แยกดำเนินคดี (เช่น ยาบ้า 12 เม็ดของแฟนสาวที่แยกดำเนินคดี) มารวมเป็นของกลางของผู้ต้องหารายนี้เด็ดขาด\n"
        "คำเตือนพิเศษเรื่องเลขไทยและเลขบัตรประชาชน: ระวังความสับสนของฟอนต์เลขไทย เช่น ๓๕ คือ 35 (สามสิบห้า) อย่าสับสนเลข ๓ (3) เป็น ๗ (7) หรือเลขอื่น จงดูรายการของกลางทางการให้ตรงกับเอกสารจริง และแปลงเลขบัตรประชาชน 13 หลักให้ถูกต้อง"
    )

    parts = [{"text": prompt}]
    if text and text.strip():
        parts.append({"text": f"--- ข้อความเอกสาร ({filename}) ---\n" + text[:25000]})
    if image_bytes:
        b64_img = base64.b64encode(image_bytes).decode("utf-8")
        parts.append({
            "inlineData": {
                "mimeType": "image/jpeg",
                "data": b64_img
            }
        })
    if image_list:
        for img_b in image_list:
            if img_b:
                b64_img = base64.b64encode(img_b).decode("utf-8")
                parts.append({
                    "inlineData": {
                        "mimeType": "image/jpeg",
                        "data": b64_img
                    }
                })

    payload = {
        "contents": [{"role": "user", "parts": parts}],
        "generationConfig": {
            "temperature": 0.1,
            "responseMimeType": "application/json"
        }
    }

    for model_name in model_names:
        url = f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent"
        try:
            req = urllib.request.Request(
                url,
                data=json.dumps(payload).encode("utf-8"),
                headers={"Content-Type": "application/json", "x-goog-api-key": api_key},
                method="POST"
            )
            with urllib.request.urlopen(req, timeout=25) as resp:
                data = json.loads(resp.read().decode("utf-8"))
                candidates = data.get("candidates", [])
                if candidates and "content" in candidates[0]:
                    cand_parts = candidates[0]["content"].get("parts", [])
                    if cand_parts and "text" in cand_parts[0]:
                        raw_json_str = cand_parts[0]["text"].strip()
                        if raw_json_str.startswith("```json"):
                            raw_json_str = raw_json_str[7:]
                        if raw_json_str.startswith("```"):
                            raw_json_str = raw_json_str[3:]
                        if raw_json_str.endswith("```"):
                            raw_json_str = raw_json_str[:-3]
                        res_obj = json.loads(raw_json_str.strip())
                        logger.info(f"Gemini LLM ({model_name}) successfully extracted case structure from {filename}")
                        return res_obj
        except Exception as e:
            logger.warning(f"Gemini LLM call failed with {model_name} for {filename}: {e}")
            continue

    return None


def extract_tr14_data(image_bytes: bytes) -> Dict[str, Any]:
    """
    Extracts text from a Thai TR.14/ID image using Gemini Vision AI (if configured) or Tesseract OCR,
    and parses out Name, ID Card, Address, Phone Number, Role Tag, Age, and Drug Quantity.
    """
    try:
        # 1. First attempt with Google Gemini Vision AI if API key is present
        gemini_res = extract_with_gemini_vision(image_bytes)
        if gemini_res and (gemini_res.get("suspect_name") or gemini_res.get("id_card")):
            return gemini_res

        image = Image.open(io.BytesIO(image_bytes))
        clean_img = _remove_red_markup(image)
        
        # Run multi-pass OCR for high fidelity (tha-only for names + tha+eng + psm 6)
        text_tha = pytesseract.image_to_string(clean_img, lang="tha")
        text = pytesseract.image_to_string(clean_img, lang="tha+eng")
        try:
            text_psm6 = pytesseract.image_to_string(clean_img, lang="tha+eng", config="--psm 6")
        except Exception:
            text_psm6 = ""
        # Verify if text contains genuine TR.14 / Civil Registration or ID card markers
        combined_text_check = f"{text_tha} {text} {text_psm6}"
        tr14_markers = [
            "ทะเบียนราษฎร", "แบบรับรองรายการ", "ทร.14", "ท.ร.14", "สำนักทะเบียน",
            "กรมการปกครอง", "รายการเกี่ยวกับบุคคล", "บัตรประจำตัวประชาชน",
            "เลขประจำตัวประชาชน", "วันออกบัตร", "วันหมดอายุ", "สัญชาติไทย",
            "สูติบัตร", "ใบเกิด", "สำเนาทะเบียนบ้าน", "ชื่อตัวและชื่อสกุล"
        ]
        has_tr14_marker = any(k in combined_text_check for k in tr14_markers)

        extracted_data = {
            "suspect_name": None,
            "id_card": None,
            "address": None,
            "age": None,
            "phone_number": None,
            "all_phones": [],
            "role_tag": None,
            "area_tag": None,
            "drug_quantity": 0
        }
        
        THAI_CHAR = r'[\u0E00-\u0E7F]'
        PREFIXES = r'(?:(?<![ก-๙a-zA-Z])(?:นาย|หาาย|บาย|ทาย|มาย|ผาย|นางสาว|มางสาว|บา่งสาว|น\.ส\.|นาง(?!รอง)|เด็กชาย|เด็กหญิง|ด\.ช\.|ด\.ญ\.))'
        
        NON_NAME_WORDS = {
            "เพศ", "สถานภาพ", "สถานภาพบุคคล", "บุคคล", "สัญชาติ", "เกิด", "วันเดือนปี", "อายุ",
            "รหัส", "ทะเบียน", "ที่อยู่", "บิดา", "มารดา", "ผู้แจ้ง", "ประจำตัว", "ข้อมูล",
            "สอบถาม", "รายการ", "สิทธิ", "ชาย", "หญิง", "คำนำหน้านาม", "ความสัมพันธ์",
            "สัญชาติไทย", "สัญชาติเดิม", "ศาสนา", "อาชีพ", "การศึกษา", "สมรส", "โสด", "หม้าย", "หย่า", "ตาย", "รอง",
            "ข้อหา", "ผู้ต้องหา", "กล่าวหา", "ความผิด", "ฐานความผิด", "คดี", "ของกลาง", "หมายจับ", "สำเนา",
            "พฤติการณ์", "เหา", "จับกุม", "บันทึก", "ดำเนินคดี", "สืบสวน", "ปราบปราม", "พนักงานสอบสวน",
            "สถานีตำรวจ", "สภ", "บช", "บก", "ชปส", "ยาเสพติด", "ตรวจปัสสาวะ", "เมทแอมเฟตามีน", "โดยกล่าวหาว่า",
            "ร้อย", "ตชด", "ตชด215", "ฉก", "ทพ", "กอรมน", "ศปปส", "อส", "อสอ", "ภจว", "ชุดปฏิบัติการ", "กองร้อย", "หมวด", "นางรอง"
        }

        # 1. Extract 13-digit ID Card with strict keyword prefixing & barcode exclusion
        id_candidates = []
        for t_src in [text_tha, text, text_psm6]:
            if not t_src: continue
            cleaned_t = clean_thai_text(t_src)
            m_pre = re.search(r'(?:เลขประจำตัวประชาชน|เลขประจําตัวประชาชน|เลขประจำตัว|เลขประจําตัว|บัตรประจำตัวประชาชน|บัตรประจําตัวประชาชน)\s*[:\.]?\s*([0-9\s\-]{13,25})', cleaned_t)
            if m_pre:
                cand = re.sub(r'\D', '', m_pre.group(1))
                if len(cand) >= 13 and not cand.startswith('9'):
                    id_candidates.append(cand[:13])
        if not id_candidates:
            # Match standard hyphenated ID pattern 1-XXXX-XXXXX-XX-X
            for t_src in [text_tha, text, text_psm6]:
                if not t_src: continue
                cleaned_t = clean_thai_text(t_src)
                m_hyphen = re.search(r'\b([0-9])\s*-\s*([0-9]{4})\s*-\s*([0-9]{5})\s*-\s*([0-9]{2})\s*-\s*([0-9])\b', cleaned_t)
                if m_hyphen:
                    cand = "".join(m_hyphen.groups())
                    if not cand.startswith('9'):
                        id_candidates.append(cand)
        if not id_candidates:
            # Fallback
            for t_src in [text_tha, text, text_psm6]:
                if not t_src: continue
                for m in re.finditer(r'(\d)[\s-]?(\d{4})[\s-]?(\d{5})[\s-]?(\d{2})[\s-]?(\d)', t_src):
                    cand = "".join(m.groups())
                    if len(cand) == 13 and not cand.startswith('9'):
                        id_candidates.append(cand)

        # Prioritize candidate that passes checksum
        valid_id = next((c for c in id_candidates if validate_thai_id_checksum(c)), None)
        extracted_data["id_card"] = valid_id or (id_candidates[0] if id_candidates else None)

        POLICE_SIGN_FILTER = ["ตชด", "ชปส", "กอ.รมน", "อส", "ร้อย", "นางรอง", "ศป.ปส", "ฉก.ทพ", "ภ.จว", "สภ"]

        # 2. Extract Suspect Name (Handles 2-column TR.14 layouts and strict Thai characters)
        for t_src in [text_tha, text, text_psm6]:
            if not t_src or extracted_data.get("suspect_name"): continue
            name_match1 = re.search(
                r'(?:ชื่อ[\s\-_–/\"\'\.]*ชื่อสกุล|ชื่อ[\s\-_–/\"\'\.]*สกุล|ชื่อตัวและชื่อสกุล|ชื่อตัว)\s*[:\.]?\s*(' + PREFIXES + r'?)\s*(' + THAI_CHAR + r'{2,})\s+(' + THAI_CHAR + r'{2,})',
                t_src
            )
            if name_match1:
                p = name_match1.group(1).strip()
                f = name_match1.group(2).strip()
                l = name_match1.group(3).strip()
                is_bad = (
                    f in NON_NAME_WORDS or l in NON_NAME_WORDS or
                    f in ["รี", "ร้อย"] or l in ["รี", "ร้อย"] or
                    any(k in f for k in POLICE_SIGN_FILTER) or any(k in l for k in POLICE_SIGN_FILTER) or
                    any(w in f for w in ["เพศ", "สถานภาพ", "วันเดือนปี", "สัญชาติ", "ข้อหา", "ผู้ต้องหา", "กล่าวหา", "เหา"]) or
                    any(w in l for w in ["เพศ", "สถานภาพ", "วันเดือนปี", "สัญชาติ", "ข้อหา", "ผู้ต้องหา", "กล่าวหา"]) or
                    (not p and (len(f) < 3 or len(l) < 3 or f.startswith("ผู้") or l.startswith("ข้อ")))
                )
                if not is_bad:
                    extracted_data["suspect_name"] = f"{p}{f} {l}".strip()

        if not extracted_data.get("suspect_name"):
            for t_src in [text_tha, text, text_psm6]:
                if not t_src or extracted_data.get("suspect_name"): continue
                for m in re.finditer(r'(' + PREFIXES + r')\s*(' + THAI_CHAR + r'{2,})\s+(' + THAI_CHAR + r'{2,})', t_src):
                    p = m.group(1).strip()
                    f_word = m.group(2).strip()
                    l_word = m.group(3).strip()
                    if (
                        f_word not in NON_NAME_WORDS and l_word not in NON_NAME_WORDS and
                        f_word not in ["รี", "ร้อย"] and l_word not in ["รี", "ร้อย"] and
                        not any(k in f_word for k in POLICE_SIGN_FILTER) and
                        not any(k in l_word for k in POLICE_SIGN_FILTER) and
                        not any(w in f_word for w in ["เพศ", "สถานภาพ", "วันเดือนปี", "สัญชาติ", "รอง", "ข้อหา", "ผู้ต้องหา", "เหา"]) and
                        not any(w in l_word for w in ["เพศ", "สถานภาพ", "วันเดือนปี", "สัญชาติ", "รอง", "ข้อหา", "ผู้ต้องหา"])
                    ):
                        extracted_data["suspect_name"] = f"{p}{f_word} {l_word}".strip()
                        break

        if extracted_data.get("suspect_name"):
            clean_name = normalize_thai_name_candidate(extracted_data["suspect_name"])
            if clean_name and any(bad in clean_name for bad in ["เพศ", "สถานภาพบุคคล", "วันเดือนปี", "ข้อหา", "ผู้ต้องหา", "เหา"]):
                extracted_data["suspect_name"] = None
            else:
                extracted_data["suspect_name"] = clean_name

        # 3. Extract Phone Numbers & Role / Area Tags
        phones = []
        combined_text = f"{text} {text_psm6}"
        for m in re.finditer(r'\b(0[689]\d{8})\b|\b(0[689]\d[-\s]?\d{3}[-\s]?\d{4})\b', combined_text):
            clean = re.sub(r'\D', '', m.group(0))
            if len(clean) == 10 and clean not in phones:
                phones.append(clean)
        
        if phones:
            extracted_data["phone_number"] = phones[0]
            extracted_data["all_phones"] = phones

        if "คนจดเบอร์" in combined_text or "จดเบอร์" in combined_text:
            extracted_data["role_tag"] = "คนจดเบอร์ (Sim Mule)"
        elif "ซิมม้า" in combined_text:
            extracted_data["role_tag"] = "ซิมม้า"

        m_tag = re.search(r'(บร\s*\d+|สต\s*\d+|พยัค\s*\d+แขบ|แขบ\s*\d+)', combined_text)
        if m_tag:
            extracted_data["area_tag"] = m_tag.group(1).replace(" ", "")

        # 4. Multi-Tier Multi-Pass Smart Address Extraction
        addr_primary = parse_tr14_address(text)
        addr_psm6 = parse_tr14_address(text_psm6) if text_psm6 else None

        if addr_psm6 and "ต." in addr_psm6:
            extracted_data["address"] = addr_psm6
        elif addr_primary:
            extracted_data["address"] = addr_primary
        else:
            extracted_data["address"] = addr_psm6

        # 5. Extract Age
        age_match = re.search(r'(?:อายุ|Age)\s*[:\.]?\s*(\d+)\s*(?:ปี)?', text)
        if age_match:
            extracted_data["age"] = age_match.group(1)

        # 6. Extract Father & Mother (Family Link Metadata)
        family_info = {}
        m_f_cid = re.search(r'(?:เลขประจำตัวประชาชนบิดา|เลขประจำตัวบิดา|บิดา)\s*[:\.]?\s*(\d[\d\s\-]{12,17}\d)', combined_text)
        m_f_name = re.search(r'ชื่อบิดา\s*[:\.]?\s*([^\n\r]+?)(?=(?:\s+สัญชาติ|\s+เพศ|\s+เลขประจำตัว|\n|$))', combined_text)
        if m_f_cid:
            f_cid_clean = re.sub(r'\D', '', m_f_cid.group(1))
            if len(f_cid_clean) == 13:
                family_info["father_id"] = f_cid_clean
        if m_f_name:
            f_name_clean = clean_thai_text(m_f_name.group(1)).strip()
            if f_name_clean and len(f_name_clean) >= 2 and f_name_clean not in ["ไม่ระบุ", "-"]:
                family_info["father_name"] = f_name_clean

        m_m_cid = re.search(r'(?:เลขประจำตัวประชาชนมารดา|เลขประจำตัวมารดา|มารดา)\s*[:\.]?\s*(\d[\d\s\-]{12,17}\d)', combined_text)
        m_m_name = re.search(r'ชื่อมารดา\s*[:\.]?\s*([^\n\r]+?)(?=(?:\s+สัญชาติ|\s+เพศ|\s+เลขประจำตัว|\n|$))', combined_text)
        if m_m_cid:
            m_cid_clean = re.sub(r'\D', '', m_m_cid.group(1))
            if len(m_cid_clean) == 13:
                family_info["mother_id"] = m_cid_clean
        if m_m_name:
            m_name_clean = clean_thai_text(m_m_name.group(1)).strip()
            if m_name_clean and len(m_name_clean) >= 2 and m_name_clean not in ["ไม่ระบุ", "-"]:
                family_info["mother_name"] = m_name_clean

        if family_info:
            extracted_data["family_info"] = family_info

        # Strict validation: image must either have genuine TR.14/ID markers or a validated 13-digit Thai ID card
        if not has_tr14_marker and not extracted_data.get("id_card"):
            logger.info("Image does not contain genuine TR.14/ID card markers or 13-digit ID; skipping TR.14 parsing.")
            return None

        return extracted_data
    except Exception as e:
        logger.error(f"OCR extraction failed: {e}")
        return {
            "suspect_name": None,
            "id_card": None,
            "address": None,
            "age": None,
            "phone_number": None,
            "all_phones": [],
            "role_tag": None,
            "area_tag": None,
            "drug_quantity": 0
        }

def extract_bank_slip_data(image_bytes: bytes, caption: Optional[str] = None) -> Optional[Dict[str, Any]]:
    """
    Detects and extracts structured financial data from Thai mobile banking transfer slips.
    Recognizes Sender, Receiver, Bank Accounts, Amount, Date/Time, and Reference IDs.
    """
    try:
        image = Image.open(io.BytesIO(image_bytes))
        text = pytesseract.image_to_string(image, lang="tha+eng")
        try:
            text_psm6 = pytesseract.image_to_string(image, lang="tha+eng", config="--psm 6")
        except Exception:
            text_psm6 = ""
        
        full_ocr = text + "\n" + text_psm6
        full_lower = full_ocr.lower()
        
        # 1. Detection Keywords
        bank_keywords = [
            "bangkok bank", "kasikornbank", "kbank", "scb", "krungthai", "promptpay",
            "พร้อมเพย์", "รายการสําเร็จ", "รายการสำเร็จ", "โอนเงินสำเร็จ", "โอนสำเร็จ",
            "ธนาคารกรุงเทพ", "ธนาคารกสิกรไทย", "ธนาคารไทยพาณิชย์", "ธนาคารกรุงไทย",
            "ธนาคารทหารไทยธนชาต", "ttb", "ออมสิน", "mymo", "รหัสอ้างอิง", "เลขที่อ้างอิง",
            "ค่าธรรมเนียม", "จํานวนเงิน", "จำนวนเงิน", "successful", "transfer successful"
        ]
        kw_count = sum(1 for kw in bank_keywords if kw in full_lower)
        if kw_count < 2 and not ("สลิป" in (caption or "") or "โอน" in (caption or "")):
            return None
            
        # 2. Extract Names (Sender / Receiver)
        PREFIXES = r'(?:(?<![ก-๙a-zA-Z])(?:นาย|นางสาว|น\.ส\.|นาง|เด็กชาย|เด็กหญิง|ด\.ช\.|ด\.ญ\.))'
        THAI_CHAR = r'[\u0E00-\u0E7F]'
        name_pattern = re.compile(r'(' + PREFIXES + r'\s*' + THAI_CHAR + r'{2,}\s+' + THAI_CHAR + r'{2,})')
        
        all_names = []
        for line in full_ocr.split("\n"):
            line_clean = line.strip()
            m = name_pattern.search(line_clean)
            if m:
                n = clean_thai_text(m.group(1).strip())
                if n not in all_names and not any(w in n for w in ["สถานภาพ", "เพศ", "วันเดือนปี", "ประจำตัว", "เจ้าหน้าที่"]):
                    all_names.append(n)
        
        sender_name = all_names[0] if len(all_names) >= 1 else None
        receiver_name = all_names[1] if len(all_names) >= 2 else None
        
        # Check if caption specifies destination / receiver name
        if caption:
            cap_m = re.search(r'(' + PREFIXES + r'\s*' + THAI_CHAR + r'{2,}(?:\s+' + THAI_CHAR + r'{2,})?)', caption)
            if cap_m:
                cap_name = cap_m.group(1).strip()
                if not receiver_name or cap_name != sender_name:
                    if "ปลายทาง" in caption or "โอนไป" in caption or "รับ" in caption:
                        receiver_name = cap_name
        
        # 3. Extract Accounts
        acc_pattern = re.compile(r'\b([0-9]{3}[-\s]?[0-9][-\s]?[xX%*0-9]{3,6}[-\s]?[0-9]{3,4})\b')
        accounts = []
        for m in acc_pattern.finditer(full_ocr):
            acc = m.group(1).replace("%", "x").replace("*", "x").replace(" ", "-")
            if acc not in accounts:
                accounts.append(acc)
                
        sender_acc = accounts[0] if len(accounts) >= 1 else None
        receiver_acc = accounts[1] if len(accounts) >= 2 else None
        
        # 4. Extract Banks
        sender_bank = "ธนาคารกรุงเทพ" if "bangkok bank" in full_lower or "ธนาคารกรุงเทพ" in full_lower else None
        receiver_bank = "ธนาคารกรุงไทย" if "ธนาคารกรุจไทย" in full_ocr or "ธนาคารกรุงไทย" in full_ocr or "krungthai" in full_lower else sender_bank
        if not sender_bank:
            if "kbank" in full_lower or "กสิกร" in full_lower: sender_bank = "ธนาคารกสิกรไทย"
            elif "scb" in full_lower or "ไทยพาณิชย์" in full_lower: sender_bank = "ธนาคารไทยพาณิชย์"
            elif "ttb" in full_lower: sender_bank = "ธนาคารทหารไทยธนชาต"
            elif "ออมสิน" in full_lower: sender_bank = "ธนาคารออมสิน"
        
        # 5. Extract Amount
        amt_match = re.search(r'([0-9]{1,3}(?:,[0-9]{3})*\.[0-9]{2})', full_ocr)
        amount = amt_match.group(1) if amt_match else None
        
        # 6. Extract Ref
        ref_match = re.search(r'([0-9]{20,35})', full_ocr)
        ref_no = ref_match.group(1) if ref_match else None

        # 7. Extract Date & Time
        THAI_MONTHS = r'(?:ม\.ค\.|ก\.พ\.|มี\.ค\.|เม\.ย\.|พ\.ค\.|มิ\.ย\.|ก\.ค\.|ส\.ค\.|ก\.ย\.|ต\.ค\.|พ\.ย\.|ธ\.ค\.|มกราคม|กุมภาพันธ์|มีนาคม|เมษายน|พฤษภาคม|มิถุนายน|กรกฎาคม|สิงหาคม|กันยายน|ตุลาคม|พฤศจิกายน|ธันวาคม)'
        date_match = re.search(r'(\d{1,2}\s+' + THAI_MONTHS + r'\s+\d{2,4}(?:\s*,\s*|\s+)?(?:\d{1,2}:\d{2}(?::\d{2})?(?:\s*น\.)?)?)', full_ocr)
        if not date_match:
            date_match = re.search(r'(\d{1,2}[/\-\.]\d{1,2}[/\-\.]\d{2,4}(?:\s+\d{1,2}:\d{2}(?::\d{2})?)?)', full_ocr)
        
        transaction_datetime = date_match.group(1).strip() if date_match else None
        if not transaction_datetime and ref_no and len(ref_no) >= 12 and ref_no.startswith(('202', '203')):
            try:
                yr = int(ref_no[:4])
                mo = int(ref_no[4:6])
                dy = int(ref_no[6:8])
                hr = int(ref_no[8:10])
                mi = int(ref_no[10:12])
                thai_months_abbr = ["", "ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]
                if 1 <= mo <= 12 and 1 <= dy <= 31 and 0 <= hr <= 23 and 0 <= mi <= 59:
                    th_mo = thai_months_abbr[mo]
                    th_yr = (yr + 543) % 100
                    transaction_datetime = f"{dy} {th_mo} {th_yr} {hr:02d}:{mi:02d} น."
            except Exception:
                pass
        
        return {
            "is_bank_slip": True,
            "sender_name": sender_name,
            "sender_bank": sender_bank,
            "sender_account": sender_acc,
            "receiver_name": receiver_name,
            "receiver_bank": receiver_bank,
            "receiver_account": receiver_acc,
            "amount": amount,
            "ref_no": ref_no,
            "transaction_datetime": transaction_datetime,
            "notes": caption or "สลิปการโอนเงิน",
            "raw_text_snippet": text[:200]
        }
    except Exception as e:
        logger.error(f"Error in bank slip extraction: {e}")
        return None


def _extract_tables_from_docx(doc) -> List[str]:
    """
    Extracts structured text lines from all Word tables (e.g. Seized items, Suspect lists).
    Formats evidence and suspect rows into clear natural Thai sentences for seamless downstream parsing.
    """
    table_lines = []
    if not hasattr(doc, 'tables'):
        return table_lines

    for t_idx, table in enumerate(doc.tables, 1):
        headers = []
        is_evidence_table = False
        is_suspect_table = False
        col_map = {}

        for r_idx, row in enumerate(table.rows):
            # deduplicate adjacent merged cell texts
            row_cells = []
            for cell in row.cells:
                ctext = re.sub(r'\s+', ' ', cell.text).strip()
                if ctext and (not row_cells or ctext != row_cells[-1]):
                    row_cells.append(ctext)

            if not row_cells:
                continue

            joined_row = " ".join(row_cells)

            # Header detection
            if r_idx == 0:
                headers = [c.lower() for c in row_cells]
                header_str = " ".join(headers)
                if any(k in header_str for k in ["ของกลาง", "รายการ", "สิ่งของ", "อาวุธ", "ยาเสพติด"]):
                    is_evidence_table = True
                    for c_idx, h in enumerate(headers):
                        if any(k in h for k in ['ของกลาง', 'รายการ', 'สิ่งของ']): col_map['item'] = c_idx
                        elif 'จำนวน' in h: col_map['qty'] = c_idx
                        elif 'หน่วย' in h: col_map['unit'] = c_idx
                        elif any(k in h for k in ['หมายเหตุ', 'ที่พบ']): col_map['notes'] = c_idx

                if any(k in header_str for k in ["ผู้ต้องหา", "ชื่อ", "นามสกุล", "บุคคล", "ผู้ถูกจับ"]):
                    is_suspect_table = True
                    for c_idx, h in enumerate(headers):
                        if 'ชื่อ' in h: col_map['name'] = c_idx
                        elif any(k in h for k in ['บัตร', 'ประชาชน', 'id']): col_map['cid'] = c_idx
                        elif 'อายุ' in h: col_map['age'] = c_idx
                        elif any(k in h for k in ['ที่อยู่', 'บ้าน']): col_map['addr'] = c_idx
                        elif any(k in h for k in ['ข้อหา', 'พฤติการณ์']): col_map['charge'] = c_idx
                continue

            # Format row
            if is_evidence_table and ('item' in col_map or len(row_cells) >= 3):
                item_name = row_cells[col_map['item']] if 'item' in col_map and col_map['item'] < len(row_cells) else (row_cells[1] if len(row_cells) > 1 else row_cells[0])
                qty = row_cells[col_map['qty']] if 'qty' in col_map and col_map['qty'] < len(row_cells) else (row_cells[2] if len(row_cells) > 2 else '')
                unit = row_cells[col_map['unit']] if 'unit' in col_map and col_map['unit'] < len(row_cells) else (row_cells[3] if len(row_cells) > 3 else '')
                notes = row_cells[col_map['notes']] if 'notes' in col_map and col_map['notes'] < len(row_cells) else (row_cells[4] if len(row_cells) > 4 else '')
                table_lines.append(f"[บัญชีของกลาง ลำดับ {r_idx}] {item_name} จำนวน {qty} {unit} ({notes})".strip())

            elif is_suspect_table and ('name' in col_map or len(row_cells) >= 3):
                name = row_cells[col_map['name']] if 'name' in col_map and col_map['name'] < len(row_cells) else (row_cells[1] if len(row_cells) > 1 else row_cells[0])
                cid = row_cells[col_map['cid']] if 'cid' in col_map and col_map['cid'] < len(row_cells) else (row_cells[2] if len(row_cells) > 2 else '')
                age = row_cells[col_map['age']] if 'age' in col_map and col_map['age'] < len(row_cells) else (row_cells[3] if len(row_cells) > 3 else '')
                addr = row_cells[col_map['addr']] if 'addr' in col_map and col_map['addr'] < len(row_cells) else (row_cells[4] if len(row_cells) > 4 else '')
                charge = row_cells[col_map['charge']] if 'charge' in col_map and col_map['charge'] < len(row_cells) else (row_cells[5] if len(row_cells) > 5 else '')
                
                if addr and not addr.startswith('บ้านเลขที่') and not addr.startswith('อยู่'):
                    addr = f"บ้านเลขที่ {addr}"
                if age and 'ปี' not in age:
                    age = f"{age} ปี"
                table_lines.append(f"[ผู้ต้องหาที่ {r_idx}] {name} อายุ {age} อยู่{addr} เลขประจำตัวประชาชน {cid} ข้อหา {charge}".strip())

            else:
                if any(k in joined_row for k in ["ยาบ้า", "ไอซ์", "ปืน", "อาวุธ", "เงินสด"]):
                    table_lines.append(f"[ของกลาง ลำดับ {r_idx}] " + " | ".join(row_cells))
                elif any(k in joined_row for k in ["นาย", "นาง", "น.ส.", "ด.ช.", "ด.ญ."]):
                    table_lines.append(f"[ผู้ต้องหาที่ {r_idx}] " + " | ".join(row_cells))
                else:
                    table_lines.append(" | ".join(row_cells))

    return table_lines


def _extract_text_from_binary_doc(file_bytes: bytes) -> str:
    """
    Pure Python stream text extractor for legacy Word 97-2003 (.doc) binary OLE2 files.
    Recovers UTF-16LE and Thai CP874 text chunks without requiring external OS tools.
    """
    text_pieces = []
    
    # 1. Try olefile if available
    try:
        import olefile
        if olefile.isOleFile(io.BytesIO(file_bytes)):
            ole = olefile.OleFileIO(io.BytesIO(file_bytes))
            if ole.exists('WordDocument'):
                stream = ole.openstream('WordDocument').read()
                # Scan for UTF-16LE text runs (min 4 chars)
                utf16_runs = re.findall(rb'(?:[\x20-\x7e\x00-\x7e][\x00]|[\x00-\xff][\x0e]){4,}', stream)
                for run in utf16_runs:
                    try:
                        decoded = run.decode('utf-16le', errors='ignore').strip()
                        if len(decoded) >= 4 and any('\u0e00' <= c <= '\u0e7f' or c.isalnum() for c in decoded):
                            text_pieces.append(decoded)
                    except Exception:
                        pass
    except Exception as e:
        logger.debug(f"olefile extraction fallback: {e}")

    # 2. General binary regex sweep for Thai UTF-16LE
    if not text_pieces:
        for m in re.finditer(rb'((?:[\x20-\x7e][\x00]|[\x01-\x7f][\x0e]){4,})', file_bytes):
            try:
                decoded = m.group(1).decode('utf-16le', errors='ignore').strip()
                if len(decoded) >= 4 and any('\u0e00' <= c <= '\u0e7f' or c.isalnum() for c in decoded):
                    text_pieces.append(decoded)
            except Exception:
                pass

    # 3. Also check CP874 (Windows-874 Thai 8-bit encoding: 0xA1-0xFB)
    cp874_runs = re.findall(rb'[\xa1-\xfb\x20-\x7e]{6,}', file_bytes)
    for run in cp874_runs:
        try:
            decoded = run.decode('cp874', errors='ignore').strip()
            if any('\u0e00' <= c <= '\u0e7f' for c in decoded) and len(decoded) >= 5:
                text_pieces.append(decoded)
        except Exception:
            pass

    return "\n".join(text_pieces)


def extract_document_data(file_bytes: bytes, filename: str, caption: Optional[str] = None) -> Dict[str, Any]:
    """
    Extracts text from DOCX, PDF, Bank Slips, TR.14 images, or Phone list images.
    """
    ext = filename.split('.')[-1].lower() if '.' in filename else ''
    full_text = ""
    pdf_key_images: List[bytes] = []
    
    try:
        if ext in ['docx', 'doc']:
            # 1. Check if it's a modern OpenXML (.docx) even if named .doc (starts with PK)
            if file_bytes.startswith(b'PK') and docx is not None:
                doc = docx.Document(io.BytesIO(file_bytes))
                paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
                table_lines = _extract_tables_from_docx(doc)
                full_text = "\n".join(paragraphs + table_lines)
            else:
                # 2. Legacy binary OLE2 Word 97-2003 (.doc)
                full_text = _extract_text_from_binary_doc(file_bytes)
                if not full_text.strip() and docx is not None:
                    try:
                        doc = docx.Document(io.BytesIO(file_bytes))
                        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
                        table_lines = _extract_tables_from_docx(doc)
                        full_text = "\n".join(paragraphs + table_lines)
                    except Exception:
                        pass
        elif ext == 'pdf':
            # 1. Try extracting direct text via pypdf first
            pdf_text_direct = ""
            try:
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                pages_text = [p.extract_text() for p in reader.pages if p.extract_text()]
                candidate_pdf_text = "\n".join(pages_text)
                if is_gibberish_text(candidate_pdf_text):
                    logger.info(f"pypdf direct text from {filename} is corrupt font encoding (gibberish); falling back to clean image OCR.")
                    pdf_text_direct = ""
                else:
                    pdf_text_direct = candidate_pdf_text
            except Exception as e:
                logger.info(f"pypdf direct extraction failed for {filename}: {e}")

            # 2. Convert to images for OCR page-by-page to keep RAM minimal (~12MB per page instead of 600MB+)
            text_parts = []
            try:
                total_pages = 1
                if pdfinfo_from_bytes:
                    try:
                        info = pdfinfo_from_bytes(file_bytes)
                        total_pages = int(info.get("Pages", 1))
                    except Exception:
                        total_pages = 1
                for page_idx in range(1, total_pages + 1):
                    page_images = convert_from_bytes(file_bytes, fmt='jpeg', first_page=page_idx, last_page=page_idx) if convert_from_bytes else []
                    for img in page_images:
                        page_text = pytesseract.image_to_string(img, lang="tha+eng") if pytesseract else ""
                        text_parts.append(page_text)
                        
                        # Retain key page images for Gemini Vision
                        has_contraband_kw = any(
                            k in page_text for k in [
                                "ของกลาง", "ตรวจนับ", "ยาบ้า", "ไอซ์", "อาวุธ", "ปืน",
                                "เม็ด", "กรัม", "บาท"
                            ]
                        )
                        is_key_page = (total_pages <= 16) or (page_idx in [1, 2]) or has_contraband_kw
                        if is_key_page and len(pdf_key_images) < 16:
                            buf = io.BytesIO()
                            w, h = img.size
                            if max(w, h) > 1600:
                                scale = 1600.0 / max(w, h)
                                img_small = img.resize((int(w * scale), int(h * scale)), Image.Resampling.LANCZOS)
                                img_small.save(buf, format="JPEG", quality=80)
                            else:
                                img.save(buf, format="JPEG", quality=80)
                            pdf_key_images.append(buf.getvalue())
                    del page_images
            except Exception as ocr_err:
                logger.warning(f"Page-by-page OCR failed on {filename}, falling back: {ocr_err}")
                if not text_parts and convert_from_bytes and pytesseract:
                    images = convert_from_bytes(file_bytes, fmt='jpeg')
                    for p_i, img in enumerate(images, 1):
                        p_t = pytesseract.image_to_string(img, lang="tha+eng")
                        text_parts.append(p_t)
                        if (p_i in [1, 2] or len(images) <= 16) and len(pdf_key_images) < 16:
                            buf = io.BytesIO()
                            img.save(buf, format="JPEG", quality=80)
                            pdf_key_images.append(buf.getvalue())
            ocr_text_raw = "\n".join([t.strip() for t in text_parts if t.strip()])

            combined_check = f"{pdf_text_direct}\n{ocr_text_raw}"
            pages_clean = [clean_thai_text(t) for t in text_parts]
            full_clean = "\n".join(pages_clean)

            # -----------------------------------------------------------------
            # 2.0 Check if PDF is a Bank Statement (รายการเดินบัญชี)
            # -----------------------------------------------------------------
            STATEMENT_KEYWORDS = [
                "รายการเดินบัญชี", "statement", "ยอดคงเหลือ", "ยอดเงินคงเหลือ",
                "เดบิต", "เครดิต", "เลขที่บัญชี", "ยอดเงินฝาก", "ยอดเงินถอน"
            ]
            is_stmt_doc = sum(1 for kw in STATEMENT_KEYWORDS if kw in combined_check.lower()) >= 2
            if is_stmt_doc and not any(k in combined_check for k in ["หมายจับ", "ทร.14", "ข้อมูลทะเบียนราษฎร", "ในพระปรมาภิไธย"]):
                logger.info(f"Detected Bank Statement PDF: {filename}")
                return {
                    "suspects": [],
                    "behavior": f"เอกสาร Statement/รายการเดินบัญชี ({filename})",
                    "is_statement": True,
                    "raw_text": full_clean[:3000],
                    "ocr_raw_text": full_clean
                }

            # -----------------------------------------------------------------
            # 2.1 Check if PDF is an Arrest Warrant (หมายจับ) or Suspect Profile
            # -----------------------------------------------------------------
            is_warrant_doc = any(k in combined_check for k in [
                "หมายจับ", "ในพระปรมาภิไธยพระมหากษัตริย์",
                "ตําหนิรูปพรรณผู้กระทําผิด", "ตำหนิรูปพรรณผู้กระทำผิด", "ตําหนิรูปพรรณ", "ตำหนิรูปพรรณ"
            ])

            if is_warrant_doc:
                logger.info(f"Detected Arrest Warrant / Suspect Profile PDF: {filename}")
                PREFIXES = r'(?:(?<![ก-๙a-zA-Z])(?:นาย|หาาย|บาย|ทาย|มาย|ผาย|นางสาว|มางสาว|บา่งสาว|น\.ส\.|นาง(?!รอง)|เด็กชาย|เด็กหญิง|ด\.ช\.|ด\.ญ\.))'

                # Primary Name from warrant command or description header
                primary_name = None
                m_target = re.search(r'(?:เพราะฉะนั้นให้ท่านจับตัว|ให้ท่านจับตัว|ด้วย)\s*([^\n\r]+?)(?=(?:\s+ซึ่งต้องหา|\s+หมายเลขประจำตัว|\n|$))', full_clean)
                if m_target:
                    cand = m_target.group(1).strip()
                    m_pfx = re.search(PREFIXES + r'\s*[^\s\(]+(?:\s+[^\s\(]+)?', cand)
                    if m_pfx:
                        primary_name = m_pfx.group(0).strip()

                if not primary_name:
                    m_tname = re.search(r'(?:นามสกุล\s*\(\s*ภาษาไทย\s*\)|ชื่อ-ชื่อสกุล|ชื่อตัว-ชื่อสกุล|ชื่อ)\s*[:\.]?\s*([^\n\r]+?)(?=(?:\s+เพศ|\s+วันเดือนปีเกิด|\n|$))', full_clean)
                    if m_tname:
                        cand = m_tname.group(1).strip()
                        m_pfx = re.search(PREFIXES + r'\s*[^\s\(]+(?:\s+[^\s\(]+)?', cand)
                        if m_pfx:
                            primary_name = m_pfx.group(0).strip()

                # Clean name typos and strip trailing punctuation
                if primary_name:
                    primary_name = normalize_thai_name_candidate(primary_name)

                # Primary CID: Look for valid 13-digit ID across standard and psm 6 OCR
                cids = []
                # Check each page with standard and psm 6 for clear ID extraction
                for img in images:
                    for cfg in ["", "--psm 6"]:
                        try:
                            t_page = clean_thai_text(pytesseract.image_to_string(img, lang="tha+eng", config=cfg) if cfg else pytesseract.image_to_string(img, lang="tha+eng"))
                            # 1. Hyphenated pattern 1-XXXX-XXXXX-XX-X
                            for m in re.finditer(r'\b(\d)[\s-]?(\d{4})[\s-]?(\d{5})[\s-]?(\d{2})[\s-]?(\d)\b', t_page):
                                cid_clean = ''.join(m.groups())
                                if len(cid_clean) == 13 and not cid_clean.startswith('9') and not cid_clean.startswith('000'):
                                    if cid_clean not in cids:
                                        cids.append(cid_clean)
                            # 2. Prefixed keyword pattern
                            m_pre = re.search(r'(?:เลขประจำตัวประชาชน|เลขประจําตัวประชาชน|บัตรประจำตัวประชาชน)\s*[:\.]?\s*([0-9\s\-]{13,25})', t_page)
                            if m_pre:
                                cand = re.sub(r'\D', '', m_pre.group(1))
                                if len(cand) == 13 and not cand.startswith('9') and not cand.startswith('000'):
                                    if cand not in cids:
                                        cids.append(cand)
                        except Exception:
                            pass

                primary_cid = None
                # If TR.14 page is attached, prioritize the CID on that page
                for p in pages_clean:
                    if any(k in p for k in ["ขอมูลทะเบยนราษฎร", "ข้อมูลทะเบียนราษฎร", "แบบรับรองรายการทะเบียนราษฎร", "ทร.14"]):
                        for c in cids:
                            if c in p.replace('-', '').replace(' ', '') and validate_thai_id_checksum(c):
                                primary_cid = c
                                break
                if not primary_cid and cids:
                    primary_cid = next((c for c in cids if validate_thai_id_checksum(c)), None)

                # Primary Address: Prioritize TR.14 page, else warrant address
                primary_addr = None
                for p in pages_clean:
                    if any(k in p for k in ["ขอมูลทะเบยนราษฎร", "ข้อมูลทะเบียนราษฎร", "แบบรับรองรายการทะเบียนราษฎร", "ทร.14"]):
                        primary_addr = parse_tr14_address(p) or clean_and_format_thai_address(p)
                        if primary_addr:
                            break
                if not primary_addr:
                    primary_addr = clean_and_format_thai_address(full_clean)

                # Charge (ข้อหา / ฐานความผิด)
                m_chg = re.search(r'(?:ซึ่งต้องหาว่ากระทำความผิดฐาน|ฐานความผิด|ข้อหา)\s*[:\.]?\s*([^\n\r]+(?:\n[^\n\r]+)?)', full_clean)
                charge = m_chg.group(1).strip() if m_chg else None

                # Family metadata (Father, Mother, Spouse) - strictly kept as metadata, never as suspects!
                family_info = {}
                m_dad = re.search(r'(?:[ชื่อซือซื้อ]+(?:นามสกุล)?\s*บิดา|บิดาชื่อ|บิดา)\s*[:\.]?\s*(' + PREFIXES + r'\s*[^\s\(]+(?:\s+[^\s\(]+)?)', full_clean)
                if m_dad:
                    family_info["father"] = m_dad.group(1).strip()
                m_mom = re.search(r'(?:[ชื่อซือซื้อ]+(?:นามสกุล)?\s*มารด[า\']?|มารดาชื่อ|มารดา)\s*[:\.]?\s*(' + PREFIXES + r'\s*[^\s\(]+(?:\s+[^\s\(]+)?)', full_clean)
                if m_mom:
                    family_info["mother"] = m_mom.group(1).strip()

                # Age
                m_age = re.search(r'อายุ\s*[:\.]?\s*(\d{1,3})\s*ปี', full_clean)
                s_age = m_age.group(1) if m_age else None

                if primary_name or primary_cid:
                    seized_w = extract_seized_items(full_clean)
                    plates_w = extract_vehicle_plates(full_clean)
                    social_w = extract_social_media(full_clean)
                    case_w = extract_case_info(full_clean)
                    return {
                        "suspects": [{
                            "suspect_name": primary_name or "ไม่ระบุชื่อ",
                            "id_card": primary_cid,
                            "phone_number": None,
                            "age": s_age,
                            "address": primary_addr,
                            "role_tag": "ผู้ต้องหาตามหมายจับ",
                            "family_info": family_info,
                            "drug_quantity": seized_w.get("total_drug_quantity", 0)
                        }],
                        "behavior": f"หมายจับ/ตำหนิรูปพรรณ: {charge or 'คดีตามหมายจับ'} ({primary_name or ''})",
                        "seized_items": seized_w,
                        "vehicle_plates": plates_w,
                        "social_media": social_w,
                        "case_number": case_w.get("case_number"),
                        "arrest_date": case_w.get("arrest_date"),
                        "ocr_raw_text": full_clean
                    }

            # -----------------------------------------------------------------
            # 2.2 Check if PDF is a standalone TR.14 Civil Registration document
            # -----------------------------------------------------------------
            if any(k in combined_check for k in ["ข้อมูลทะเบียนราษฎร", "แบบรับรองรายการทะเบียนราษฎร", "ทร.14", "สำนักทะเบียนกลาง"]):
                tr14_cand_text = pdf_text_direct if "เลขประจำตัวประชาชน" in pdf_text_direct else ocr_text_raw
                parsed_addr = parse_tr14_address(tr14_cand_text) or clean_and_format_thai_address(tr14_cand_text)
                
                m_cid = re.search(r'(?:เลขประจำตัวประชาชน|บัตรประจำตัวประชาชน)?\s*[:\.]?\s*(\d[\d\s\-]{13,17}\d)', tr14_cand_text)
                clean_cid = re.sub(r'\D', '', m_cid.group(1)) if m_cid else None
                
                m_name = re.search(r'(?:ชื่อ-ชื่อสกุล|ชื่อตัว-ชื่อสกุล|ชื่อ)\s*[:\.]?\s*([^\n\r]+?)(?=(?:\s+เพศ|\s+วันเดือนปีเกิด|\n|$))', tr14_cand_text)
                s_name = m_name.group(1).strip() if m_name else None
                if not s_name:
                    PREFIXES = r'(?:(?<![ก-๙a-zA-Z])(?:นาย|นางสาว|น\.ส\.|นาง|เด็กชาย|เด็กหญิง|ด\.ช\.|ด\.ญ\.))'
                    m_name2 = re.search(r'(' + PREFIXES + r'[^\s\(]+(?:\s+[^\s\(]+)?)', tr14_cand_text)
                    s_name = m_name2.group(1).strip() if m_name2 else None

                m_age = re.search(r'อายุ\s*[:\.]?\s*(\d+)\s*ปี', tr14_cand_text)
                s_age = m_age.group(1) if m_age else None

                if s_name or clean_cid:
                    return {
                        "suspects": [{
                            "suspect_name": s_name,
                            "id_card": clean_cid,
                            "phone_number": None,
                            "age": s_age,
                            "address": parsed_addr,
                            "role_tag": "บุคคลเป้าหมาย (ทร.14)"
                        }],
                        "behavior": f"เอกสาร ทร.14 {s_name or ''}",
                        "ocr_raw_text": tr14_cand_text
                    }

            # Use direct text if rich in suspects, else OCR
            full_text = pdf_text_direct if len(pdf_text_direct) > 100 and "นาย" in pdf_text_direct else ocr_text_raw
        elif ext in ['jpg', 'jpeg', 'png', 'webp']:
            # 1. Check if image is a Phone List
            phone_list = extract_phone_list_data(file_bytes)
            if len(phone_list) >= 2:
                logger.info(f"Detected phone network list image with {len(phone_list)} entries")
                return {"suspects": phone_list, "behavior": f"ตรวจพบบัญชีรายชื่อเบอร์โทรศัพท์เครือข่าย {len(phone_list)} หมายเลข", "is_phone_list": True}
            
            # 2. Check if image is a Bank Transfer Slip
            bank_slip_data = extract_bank_slip_data(file_bytes, caption=caption)
            if bank_slip_data and bank_slip_data.get("is_bank_slip"):
                s_info = f"ผู้โอน: {bank_slip_data.get('sender_name') or 'ไม่ระบุ'}"
                r_info = f"บัญชีปลายทาง: {bank_slip_data.get('receiver_name') or 'ไม่ระบุ'}"
                amt_str = f" ({bank_slip_data.get('amount')} บาท)" if bank_slip_data.get('amount') else ""
                logger.info(f"Detected bank transfer slip: {s_info} -> {r_info}{amt_str}")
                return {
                    "is_bank_slip": True,
                    "bank_slip_data": bank_slip_data,
                    "suspects": [],
                    "behavior": f"หลักฐานสลิปโอนเงิน: {s_info} -> {r_info}{amt_str}"
                }

            # 3. Parse as TR.14 Civil Registration Image
            tr14_data = extract_tr14_data(file_bytes)
            if tr14_data and (tr14_data.get("suspect_name") or tr14_data.get("id_card")):
                return {"suspects": [tr14_data], "behavior": f"เอกสาร ทร.14 {tr14_data.get('suspect_name') or ''} ({tr14_data.get('role_tag') or 'บุคคลเป้าหมาย'})"}

            # 4. Check if image is a Vehicle Registration Document (DLT / เล่มทะเบียน)
            veh_doc = extract_vehicle_registration_doc(file_bytes)
            if veh_doc and veh_doc.get("is_vehicle_doc"):
                logger.info(f"Detected vehicle registration document: {veh_doc.get('plate')} {veh_doc.get('province') or ''}")
                suspects_list = []
                if veh_doc.get("owner_name") and not veh_doc.get("is_corporate_owner"):
                    suspects_list.append({
                        "suspect_name": veh_doc["owner_name"],
                        "id_card": veh_doc.get("owner_id"),
                        "phone_number": None,
                        "role_tag": "เจ้าของ/ผู้ครอบครองรถ",
                        "drug_quantity": 0
                    })
                return {
                    "is_vehicle_doc": True,
                    "vehicle_data": veh_doc,
                    "suspects": suspects_list,
                    "behavior": f"ใบคู่มือจดทะเบียนรถ ทะเบียน {veh_doc.get('plate') or '-'} {veh_doc.get('province') or ''} ({veh_doc.get('vehicle_type') or 'รถยนต์'})",
                    "ocr_raw_text": veh_doc.get("raw_text", "")
                }
            
            # 5. General Image OCR (Evidence, Mugshot, Document, Seized items, Vehicles, Social media)
            try:
                raw_img = Image.open(io.BytesIO(file_bytes))
                clean_img = _remove_red_markup(raw_img)
                img_ocr = clean_thai_text(pytesseract.image_to_string(clean_img, lang="tha+eng"))
                if len(img_ocr.strip()) >= 5:
                    full_text = img_ocr
                elif _get_gemini_api_key():
                    # Smart Fallback: Blurry/handwritten image sent directly to Gemini
                    logger.info(f"Local OCR text too short ({len(img_ocr.strip())} chars), delegating image to Gemini: {filename}")
                    gem_case = extract_case_with_gemini_llm(image_bytes=file_bytes, filename=filename)
                    if gem_case:
                        suspects_out = []
                        for s in gem_case.get("suspects", []):
                            if s.get("name"):
                                suspects_out.append({
                                    "suspect_name": s.get("name"),
                                    "id_card": s.get("id_card"),
                                    "phone_number": None,
                                    "age": str(s.get("age")) if s.get("age") else None,
                                    "address": s.get("address"),
                                    "role_tag": s.get("role") or "ผู้ต้องหา",
                                    "drug_quantity": 0
                                })
                        return {
                            "suspects": suspects_out,
                            "behavior": gem_case.get("behavior") or f"เอกสารบันทึกจับกุม {filename}",
                            "seized_items": gem_case.get("seized_items", {}),
                            "district": gem_case.get("crime_scene", {}).get("district"),
                            "subdistrict": gem_case.get("crime_scene", {}).get("subdistrict"),
                            "incident_location": gem_case.get("crime_scene", {}).get("location"),
                            "case_number": gem_case.get("case_number"),
                            "ocr_raw_text": f"[Gemini Vision AI Analysis: {filename}]"
                        }
                    return {"suspects": [], "behavior": None, "ocr_raw_text": ""}
                else:
                    return {"suspects": [], "behavior": None, "ocr_raw_text": ""}
            except Exception as ocr_err:
                logger.info(f"General image OCR error: {ocr_err}")
                return {"suspects": [], "behavior": None, "ocr_raw_text": ""}
        elif ext in ['txt', 'text', 'csv', 'log']:
            try:
                full_text = file_bytes.decode('utf-8', errors='ignore')
            except Exception:
                full_text = file_bytes.decode('cp874', errors='ignore')
        else:
            return {"suspects": [], "behavior": None}
            
        logger.info(f"Successfully extracted {len(full_text)} characters from {filename}")
        
        # Clean up multi-spaces and normalize Thai text
        raw_lines_text = clean_thai_text(full_text)
        full_text = re.sub(r'\s+', ' ', raw_lines_text)
        
        result = {
            "suspects": [],
            "behavior": None
        }
        
        PREFIXES = r'(?:(?<![ก-๙a-zA-Z])(?:นาย|หาาย|บาย|ทาย|มาย|ผาย|นางสาว|มางสาว|บา่งสาว|น\.ส\.|นาง(?!รอง)|เด็กชาย|เด็กหญิง|ด\.ช\.|ด\.ญ\.))'
        NON_SUSPECT_WORDS = [
            'สืบสวน', 'เครือข่าย', 'จัดทำ', 'บันทึก', 'รายงาน', 'บิดา', 'มารดา', 'สามี', 'ภรรยา',
            'คู่สมรส', 'ญาติ', 'พยาน', 'เจ้าหน้าที่', 'ผู้พิมพ์', 'ร.ต.อ.', 'พ.ต.ท.', 'ผู้กำกับการ',
            'พนักงานสอบสวน', 'สถานีตำรวจ', 'ผู้จับกุม', 'ผู้สั่งการ', 'ศาล', 'ตำรวจ',
            'พล.ต.อ.', 'พล.ต.ท.', 'พล.ต.ต.', 'พ.ต.อ.', 'พ.ต.ต.',
            'ร.ต.ท.', 'ร.ต.ต.', 'ด.ต.', 'จ.ส.ต.', 'ส.ต.อ.', 'ส.ต.ท.', 'ส.ต.ต.',
            'นายอำเภอ', 'ปลัดอำเภอ', 'อำเภอ', 'ผบ.ร้อย', 'ผบก.', 'รอง ผบก.', 'ผกก.',
            'ชปส.', 'ชป.ปส.', 'สขว.', 'กอ.รมน.', 'ตชด.', 'ภ.จว.',
            'ร้อย', 'กองร้อย', 'ร้อย.', 'อส.', 'อส.อ.', 'ศป.ปส.', 'ฉก.ทพ.', 'สภ.', 'บช.', 'บก.',
            'ศป.ปส.อ.', 'ร้อย.อส.', 'ฉก.ทพ.26', 'ภ.จว.บร.', 'นางรอง',
            'ไม่ทราบชื่อ', 'ชายไทย', 'หญิงไทย', 'วัยรุ่น', 'ทั้งสองราย', 'ทั้งสามราย', 'ไม่ปรากฏชื่อ', 'กลุ่มชาย'
        ]
        NARRATIVE_WORDS = [
            'เข้ามา', 'ที่บ้าน', 'ขอซื้อ', 'ยาบ้า', 'ของกลาง', 'ให้การ', 'แจ้งว่า', 'พบว่า',
            'ร่วมกัน', 'จำหน่าย', 'ตรวจค้น', 'ของตน', 'กับตน', 'ไปหาตน', 'ต่อมา', 'นำพา',
            'จับกุม', 'ส่งมอบ', 'จำหน่ายยา', 'เสพยา', 'ครอบครอง', 'ลูกพี่ลูกน้อง', 'ติดกัน',
            'พร้อม', 'พร้อมด้วย', 'รวม', 'หลบหนี', 'ตรวจยึด'
        ]
        seen_names: Dict[str, Dict[str, Any]] = {}
        seen_cids: Dict[str, Dict[str, Any]] = {}

        COMMANDER_OFFICER_KEYWORDS = [
            'โดยการอำนวยการ', 'อำนวยการ', 'ผบก.', 'ผกก.', 'นายอำเภอ', 'ปลัดอำเภอ',
            'ผบ.ร้อย', 'ประกอบด้วย', 'เจ้าหน้าที่ตำรวจ', 'เจ้าหน้าที่ทหาร', 'ผู้จับกุม', 'ชุดจับกุม'
        ]

        def _is_commander_or_officer_context(text: str, match_pos: int, span: int = 70) -> bool:
            """Check if text immediately preceding or succeeding match_pos belongs to commanders/officers."""
            pre_ctx = text[max(0, match_pos - span):match_pos]
            post_ctx = text[match_pos:min(len(text), match_pos + span)]
            return any(k in pre_ctx for k in COMMANDER_OFFICER_KEYWORDS) or any(k in post_ctx for k in COMMANDER_OFFICER_KEYWORDS)

        def _clean_suspect_name_tokens(name_str: Optional[str]) -> Optional[str]:
            if not name_str:
                return None
            cleaned = re.sub(r'(?:[\s\(]+)?(?:หรือ|ชื่อเล่น)\s*[^\s\)]+[\)]?', ' ', name_str)
            cleaned = re.sub(r'\s+', ' ', cleaned).strip()
            return cleaned if cleaned else None

        seen_names: Dict[str, Dict[str, Any]] = {}
        seen_cids: Dict[str, Dict[str, Any]] = {}

        def _add_or_update_suspect(
            name: str,
            cid: Optional[str] = None,
            age: Optional[str] = None,
            addr: Optional[str] = None,
            role: Optional[str] = None,
            require_explicit: bool = True,
            trusted_source: bool = False
        ):
            clean_name = normalize_thai_name_candidate(name)
            if not clean_name or any(k in clean_name for k in NON_SUSPECT_WORDS) or any(w in clean_name for w in NARRATIVE_WORDS):
                return
            # Reject trailing abbreviation or narrative connectives e.g. "นายณัฐพงษ์หรือแบงค์ฯ", "นาย... และนาย..."
            if clean_name.endswith('ฯ') or ' และ' in clean_name or ' กับ' in clean_name or ' นำเจ้าหน้าที่' in clean_name:
                return
            clean_sans_pfx = re.sub(r'^(?:นาย|นางสาว|น\.ส\.|นาง|เด็กชาย|เด็กหญิง|ด\.ช\.|ด\.ญ\.)\s*', '', clean_name).strip()
            # Enforce length limit: real Thai names do not exceed 40 chars
            if len(clean_sans_pfx) < 3 or clean_sans_pfx.isdigit() or len(clean_sans_pfx) > 40:
                return

            clean_cid = re.sub(r'\D', '', cid) if cid else None
            if clean_cid and len(clean_cid) == 13:
                clean_cid = repair_thai_id_checksum(clean_cid)
            if clean_cid and (len(clean_cid) != 13 or clean_cid.startswith('9') or clean_cid.startswith('000')):
                clean_cid = None

            # Policy: Only extract suspects who have 13-digit CID OR explicit numbered/warrant role indicator with address or age
            is_explicit_role = bool(role and any(k in role for k in ["ผู้ต้องหาที่", "ผู้ต้องหาลำดับ", "ผู้ต้องหาตามหมายจับ", "บุคคลเป้าหมาย"]))
            if require_explicit and not clean_cid and not is_explicit_role:
                return
            # If no CID and role is numbered, require at least address or age to prevent stray text matches in narrative
            if require_explicit and not clean_cid and is_explicit_role and not addr and not age:
                return

            formatted_addr = clean_and_format_thai_address(addr) if addr else None
            clean_display_name = _clean_suspect_name_tokens(clean_name) or clean_name

            existing = None
            if clean_cid and clean_cid in seen_cids:
                existing = seen_cids[clean_cid]
            elif clean_name in seen_names:
                existing = seen_names[clean_name]
            elif clean_display_name in seen_names:
                existing = seen_names[clean_display_name]
            else:
                # Merge if same first name when one entry has no surname or one is prefix of another
                first_name_clean = clean_sans_pfx.split()[0] if clean_sans_pfx else ""
                if first_name_clean and len(first_name_clean) >= 3:
                    for s_k, s_v in list(seen_names.items()):
                        s_sans = re.sub(r'^(?:นาย|นางสาว|น\.ส\.|นาง|เด็กชาย|เด็กหญิง|ด\.ช\.|ด\.ญ\.)\s*', '', s_k).strip()
                        s_first = s_sans.split()[0] if s_sans else ""
                        if first_name_clean == s_first:
                            if len(clean_sans_pfx.split()) == 1 or len(s_sans.split()) == 1 or clean_sans_pfx.startswith(s_sans) or s_sans.startswith(clean_sans_pfx):
                                existing = s_v
                                break

            if existing:
                old_name = existing.get("suspect_name") or ""
                old_sans = re.sub(r'^(?:นาย|นางสาว|น\.ส\.|นาง|เด็กชาย|เด็กหญิง|ด\.ช\.|ด\.ญ\.)\s*', '', old_name).strip()
                if (trusted_source or len(clean_sans_pfx.split()) > len(old_sans.split())) and clean_display_name and clean_display_name != old_name:
                    logger.info(f"Updating suspect name with more complete/trusted name: '{old_name}' -> '{clean_display_name}' (CID: {clean_cid})")
                    existing["suspect_name"] = clean_display_name
                    if old_name and old_name in seen_names:
                        del seen_names[old_name]
                    seen_names[clean_display_name] = existing

                existing_cid = existing.get("id_card")
                if clean_cid:
                    should_update_cid = (
                        not existing_cid
                        or (validate_thai_id_checksum(clean_cid) and not validate_thai_id_checksum(existing_cid))
                        or (trusted_source and validate_thai_id_checksum(clean_cid))
                    )
                    if should_update_cid:
                        if existing_cid and existing_cid in seen_cids:
                            del seen_cids[existing_cid]
                        existing["id_card"] = clean_cid
                        seen_cids[clean_cid] = existing

                if age and (trusted_source or not existing.get("age")):
                    existing["age"] = str(age).replace('ปี', '').strip()
                if formatted_addr and (trusted_source or not existing.get("address")):
                    existing["address"] = formatted_addr
                if (not existing.get("role_tag") or existing.get("role_tag") == "ผู้ต้องหา") and role:
                    existing["role_tag"] = role
            else:
                new_s = {
                    "suspect_name": clean_display_name[:40].strip(),
                    "age": str(age).replace('ปี', '').strip() if age else None,
                    "address": formatted_addr,
                    "id_card": clean_cid,
                    "phone_number": None,
                    "role_tag": role or "ผู้ต้องหา",
                    "drug_quantity": 0
                }
                result["suspects"].append(new_s)
                seen_names[clean_display_name] = new_s
                if clean_cid:
                    seen_cids[clean_cid] = new_s

        # 1.0 Prioritize official suspect section (between 'ได้ร่วมกันจับกุม' and 'พร้อมด้วยของกลาง')
        m_suspects_sec = re.search(
            r'(?:ได้ร่วมกันจับกุม|ร่วมกันจับกุม|จับกุมตัว|ร่วมกันทำการจับกุม)(.*?)(?:พร้อมด้วยของกลาง|โดยกล่าวหาว่า|ข้อหา|ของกลาง)',
            raw_lines_text,
            re.DOTALL | re.IGNORECASE
        )
        suspects_text = m_suspects_sec.group(1) if m_suspects_sec else raw_lines_text

        # 1.1 Multi-suspect extraction: Numbered suspects in suspects section
        numbered_pattern = re.finditer(
            r'(?:(?:ผู้ต้องหาที่\s*([0-9]+)|\[ผู้ต้องหา(?:ที่|\s*ลำดับ)?\s*([0-9]+)\]|(?:^|\n)\s*([0-9]+)\.\s*)(?:[^\n\r]*?))'
            r'(' + PREFIXES + r'\s*[^\s\(]+(?:\s+[^\s\(]+)?(?:(?:\s+หรือ|\s+ชื่อเล่น)[^\s\(]+)?)\s*'
            r'(?:(?:อายุ|อายุประมาณ)\s*([\d]{1,3})\s*ปี)?\s*'
            r'(?:(?:อยู่บ้านเลขที่|บ้านเลขที่|บ้านเลขที|ที่อยู่)\s*([^\n\r]+?)(?=(?:เลขประจำตัวประชาชน|บัตรประจำตัว|โทรศัพท์|อายุ|\n|\[|$)))?\s*'
            r'(?:(?:เลขประจำตัวประชาชน|บัตรประจำตัวประชาชน)\s*[:\.]?\s*([\d\s\-]+))?',
            suspects_text,
            re.IGNORECASE
        )
        for m_num in numbered_pattern:
            idx1, idx2, idx3, name_n, age_n, addr_n, cid_n = m_num.groups()
            num_str = idx1 or idx2 or idx3 or ""
            role_str = f"ผู้ต้องหาที่ {num_str}" if num_str else "ผู้ต้องหา"
            _add_or_update_suspect(name_n, cid=cid_n, age=age_n, addr=addr_n, role=role_str, require_explicit=True, trusted_source=True)

        # 1.1.1 Single or unnumbered suspects directly in official suspects section
        if not result["suspects"] and m_suspects_sec:
            single_pattern = re.finditer(
                r'(' + PREFIXES + r'\s*[^\s\(]+(?:\s+[^\s\(]+)?(?:(?:\s+หรือ|\s+ชื่อเล่น)[^\s\(]+)?)\s*'
                r'(?:(?:อายุ|อายุประมาณ)\s*([\d]{1,3})\s*ปี)?\s*'
                r'(?:(?:อยู่บ้านเลขที่|บ้านเลขที่|บ้านเลขที|ที่อยู่)\s*([^\n\r]+?)(?=(?:เลขประจำตัวประชาชน|บัตรประจำตัว|โทรศัพท์|อายุ|\n|\[|$)))?\s*'
                r'(?:(?:เลขประจำตัวประชาชน|บัตรประจำตัวประชาชน)\s*[:\.]?\s*([\d\s\-]+))?',
                suspects_text,
                re.IGNORECASE
            )
            for m_s in single_pattern:
                name_s, age_s, addr_s, cid_s = m_s.groups()
                _add_or_update_suspect(name_s, cid=cid_s, age=age_s, addr=addr_s, role="ผู้ต้องหา", require_explicit=False, trusted_source=True)

        # 1.2 Table-based suspect lines: lines starting with [ผู้ต้องหา (only if no suspects found yet)
        if not result["suspects"]:
            for line in raw_lines_text.split('\n'):
                line_str = line.strip()
                if line_str.startswith('[ผู้ต้องหา') or line_str.startswith('[บุคคล'):
                    m_tname = re.search(r'(' + PREFIXES + r'\s*[^\s\(]+(?:\s+[^\s\(]+)?)', line_str)
                    if m_tname:
                        m_tcid = re.search(r'(?:เลขประจำตัวประชาชน|บัตรประจำตัวประชาชน)?\s*[:\.]?\s*(\d[\d\s\-]{13,17}\d)', line_str)
                        s_tcid = m_tcid.group(1) if m_tcid else None
                        m_tage = re.search(r'(\d{1,3})\s*ปี', line_str)
                        s_tage = m_tage.group(1) if m_tage else None
                        s_taddr = clean_and_format_thai_address(line_str)
                        m_role = re.search(r'\[(ผู้ต้องหา(?:ที่\s*\d+)?)', line_str)
                        r_tag = m_role.group(1) if m_role else "ผู้ต้องหา"
                        _add_or_update_suspect(m_tname.group(1), cid=s_tcid, age=s_tage, addr=s_taddr, role=r_tag, require_explicit=True)

        # 1.3 Multi-suspect network chart / block extraction (only if no suspects found yet)
        if not result["suspects"]:
            block_pattern = re.finditer(
                r'(?:(?:เลขประจำตัวประชาชน|บัตรประจำตัวประชาชน)?\s*[:\.]?\s*(\d[\d\s\-]{13,17}\d)\s+)?' # Optional CID before
                r'(' + PREFIXES + r'\s*[^\s\(]+(?:\s+[^\s\(]+)?(?:(?:\s+หรือ|\s+ชื่อเล่น)[^\s\(]+)?)\s*' # Name
                r'(?:อายุ\s*([\d]{1,3})\s*ปี)?\s*' # Optional Age
                r'(?:(?:อยู่บ้านเลขที่|บ้านเลขที่|บ้านเลขที|ที่อยู่)\s*([^\n\r]+(?:\n[^\n\r]+){1,3}))?', # Localized Address Block
                raw_lines_text
            )

            for m_b in block_pattern:
                cid_pre, name_b, age_b, addr_b = m_b.groups()
                match_start = m_b.start(2)
                pre_context = raw_lines_text[max(0, match_start - 35):match_start]
                if any(w in pre_context for w in ['บิดา', 'มารดา', 'สามี', 'ภรรยา', 'ญาติ', 'ผู้พิมพ์', 'ร.ต.', 'พ.ต.', 'ผู้กำกับ']):
                    continue
                if _is_commander_or_officer_context(raw_lines_text, match_start):
                    continue
                _add_or_update_suspect(name_b, cid=cid_pre, age=age_b, addr=addr_b, role="ผู้ต้องหา", require_explicit=True)

        # 1.4 Prose report regex (requires CID, only if no suspects found yet)
        if not result["suspects"]:
            m_arrest = re.search(r'(?:ได้ร่วมกันจับกุม|ร่วมกันจับกุม|จับกุมตัว|ร่วมกันทำการจับกุม)', full_text, re.IGNORECASE)
            prose_search_text = full_text[m_arrest.start():] if m_arrest else full_text

            suspect_pattern = re.compile(
                r'(' + PREFIXES + r'\s*[^\s\(]+(?:(?:หรือ|ชื่อเล่น)[^\s\(]+)?(?:[ \t]+[^\s\(]+)?).*?'
                r'(?:อายุ\s*([\d]{1,3})\s*ปี).*?'
                r'(?:อยู่บ้านเลขที่|บ้านเลขที่|บ้านเลขที|ที่อยู่)\s*([0-9/].*?)(?=(?:เลขประจำตัวประชาชน|บัตรประจำตัว|โทรศัพท์|อายุ|\n|$)).*?'
                r'(?:เลขประจำตัวประชาชน|บัตรประจำตัวประชาชน)\s*([\d\s\-]+)',
                re.IGNORECASE
            )
            for match in suspect_pattern.finditer(prose_search_text):
                if _is_commander_or_officer_context(prose_search_text, match.start(1)):
                    continue
                _add_or_update_suspect(match.group(1), cid=match.group(4), age=match.group(2), addr=match.group(3), role="ผู้ต้องหา", require_explicit=True)

        if not result["suspects"]:
            # Fallback general extraction from prose: ONLY if 13-digit CID is present!
            m_arrest = re.search(r'(?:ได้ร่วมกันจับกุม|ร่วมกันจับกุม|จับกุมตัว|ร่วมกันทำการจับกุม)', full_text, re.IGNORECASE)
            prose_fallback_text = full_text[m_arrest.start():] if m_arrest else full_text

            s_id = None
            m_id = re.search(r'(?:เลขประจำตัวประชาชน|บัตรประจำตัวประชาชน)\s*[:\.]?\s*([\d\s\-]{13,17})', prose_fallback_text)
            if m_id:
                clean_cand = re.sub(r'[^\d]', '', m_id.group(1))
                if len(clean_cand) == 13 and not clean_cand.startswith('9') and not clean_cand.startswith('000'):
                    s_id = clean_cand
                    
            if s_id:
                s_name = None
                for m_name_cand in re.finditer(r'(' + PREFIXES + r'\s*[^\s\(]+(?:[ \t]+[^\s\(]+)?)', prose_fallback_text):
                    if _is_commander_or_officer_context(prose_fallback_text, m_name_cand.start(1)):
                        continue
                    cand_name = normalize_thai_name_candidate(m_name_cand.group(1)) or ""
                    cand_clean = re.sub(r'^(?:นาย|นางสาว|น\.ส\.|นาง|เด็กชาย|เด็กหญิง|ด\.ช\.|ด\.ญ\.)\s*', '', cand_name).strip()
                    if 3 <= len(cand_clean) <= 40 and not cand_clean.isdigit():
                        if not any(w in cand_name for w in NARRATIVE_WORDS):
                            s_name = cand_name
                            break
                
                s_addr = parse_report_address(prose_fallback_text)
                s_age = None
                m_age = re.search(r'อายุ\s*[:\.]?\s*(\d{1,3})\s*ปี', prose_fallback_text)
                if m_age:
                    s_age = m_age.group(1)
                    
                _add_or_update_suspect(s_name or "ไม่ระบุชื่อ", cid=s_id, age=s_age, addr=s_addr, role="ผู้ต้องหา", require_explicit=False)

            
        # 2. Extract Phone Number from Evidence section
        evidence_match = re.search(r'พร้อมด้วยของกลางดังนี้(.*?)พฤติการณ์', full_text)
        if evidence_match:
            evidence_text = evidence_match.group(1)
            phone_match = re.search(r'(?:เบอร์โทรศัพท์|เบอร์โทร|โทรศัพท์มือถือ|โทรศัพท์)\s*([\d\-\s]{9,12})', evidence_text)
            if phone_match and result["suspects"]:
                clean_phone = re.sub(r'[^\d]', '', phone_match.group(1))
                # Fix phone bleed: assign only to first suspect (device owner) if not already assigned
                if not result["suspects"][0].get("phone_number"):
                    result["suspects"][0]["phone_number"] = clean_phone
                    
        # 3. Extract Behavior (พฤติการณ์)
        behavior_match = re.search(r'(?:พฤติการณ์ในการจับกุม|พฤติการณ์)\s*(.*?)\s*ในการจับกุมครั้งนี้เจ้าหน้าที่ชุดจับกุมได้อธิบายให้ผู้ถูกจับทราบสิทธิ์', full_text)
        if behavior_match:
            result["behavior"] = behavior_match.group(1).strip()
            
        # 4. Extract Seized Items, Vehicle Plates, Social Media, and Case Metadata
        seized_data = extract_seized_items(full_text)
        plates_data = extract_vehicle_plates(full_text)
        social_data = extract_social_media(full_text)
        case_data = extract_case_info(full_text)

        result["seized_items"] = seized_data
        result["vehicle_plates"] = plates_data
        result["social_media"] = social_data
        result["case_number"] = case_data.get("case_number")
        result["arrest_date"] = case_data.get("arrest_date")
        result["incident_location"] = case_data.get("incident_location")
        result["district"] = case_data.get("district")
        result["subdistrict"] = case_data.get("subdistrict")
        result["ocr_raw_text"] = raw_lines_text

        # Update drug_quantity from seized items
        if seized_data.get("total_drug_quantity", 0) > 0 and result["suspects"]:
            for s in result["suspects"]:
                if not s.get("drug_quantity") or s.get("drug_quantity") == 0:
                    s["drug_quantity"] = seized_data["total_drug_quantity"]

        # 4.5 Full Quality Check & Smart Fallback via Gemini LLM
        needs_gemini_qc = False
        
        # Trigger 1: Missing suspects in a document with content
        if not result.get("suspects") and (len(full_text) > 150 or any(k in full_text for k in ["จับกุม", "ผู้ต้องหา", "ของกลาง"])):
            needs_gemini_qc = True

        # Trigger 2: Missing suspect CIDs or missing incident district in an arrest document
        if any(k in full_text for k in ["จับกุม", "ผู้ถูกจับ", "ของกลาง"]):
            if any(not s.get("id_card") for s in result.get("suspects", [])) or not result.get("district"):
                needs_gemini_qc = True
            
        # Trigger 2.5: Any suspect CID fails Thai ID checksum validation (common Tesseract OCR misreads)
        if any(s.get("id_card") and not validate_thai_id_checksum(s.get("id_card")) for s in result.get("suspects", [])):
            needs_gemini_qc = True

        # Trigger 3: Mention of contraband but zero drugs and weapons captured
        has_drug_keywords = any(k in full_text for k in ["ยาบ้า", "ไอซ์", "ยาเสพติด", "เมทแอมเฟตามีน"])
        has_weapon_keywords = any(k in full_text for k in ["อาวุธปืน", "ปืนพก", "กระสุน", "เครื่องกระสุน"])
        if (has_drug_keywords or has_weapon_keywords) and seized_data.get("total_drug_quantity", 0) == 0 and seized_data.get("weapons_count", 0) == 0:
            needs_gemini_qc = True

        # Trigger 4: Complex document with length > 1500 characters
        if len(full_text) > 1500:
            needs_gemini_qc = True

        # Trigger 5: Scanned / Photo PDF with key images available
        if ext == 'pdf' and pdf_key_images and any(k in full_text for k in ["จับกุม", "ผู้ต้องหา", "ของกลาง", "ภาพถ่าย", "ภาพประกอบ"]):
            needs_gemini_qc = True

        # Trigger 6: Blurry or very short OCR text on images
        if ext in ['jpg', 'jpeg', 'png', 'webp'] and len(full_text) < 100 and not result.get("suspects"):
            needs_gemini_qc = True

        if needs_gemini_qc and _get_gemini_api_key():
            try:
                gemini_text_input = None
                if ext == 'pdf' and not pdf_text_direct and pdf_key_images:
                    # For pure image-scanned PDFs, rely 100% on Gemini Vision without biased/corrupted local OCR text
                    gemini_text_input = None
                elif full_text and full_text.strip():
                    gemini_text_input = full_text[:25000]

                gemini_data = extract_case_with_gemini_llm(
                    text=gemini_text_input,
                    image_bytes=file_bytes if ext in ['jpg', 'jpeg', 'png', 'webp'] else None,
                    image_list=pdf_key_images if ext == 'pdf' and pdf_key_images else None,
                    filename=filename
                )
                if gemini_data:
                    has_vision = bool(ext in ['jpg', 'jpeg', 'png', 'webp'] or (ext == 'pdf' and pdf_key_images))

                    # Enrich suspects
                    for g_s in gemini_data.get("suspects", []):
                        g_name = g_s.get("name")
                        if g_name:
                            _add_or_update_suspect(
                                name=g_name,
                                cid=g_s.get("id_card"),
                                age=g_s.get("age"),
                                addr=g_s.get("address"),
                                role=g_s.get("role"),
                                trusted_source=True
                            )

                    # Enrich Seized Items
                    g_seized = gemini_data.get("seized_items", {})
                    if g_seized:
                        g_meth = int(g_seized.get("meth_pills", 0))
                        if g_meth > 0:
                            if seized_data.get("is_official_seizure") and seized_data.get("total_drug_quantity", 0) > 0:
                                # Official police seizure list in document is authoritative; keep local official total
                                pass
                            elif has_vision or seized_data.get("total_drug_quantity", 0) == 0:
                                seized_data["total_drug_quantity"] = g_meth
                                seized_data["meth_pills"] = g_meth
                            elif seized_data.get("total_drug_quantity", 0) > 1000 and g_meth < 1000:
                                # Local regex likely summed up expansion/raid quantities from other locations
                                seized_data["total_drug_quantity"] = g_meth
                                seized_data["meth_pills"] = g_meth
                            elif not seized_data.get("is_official_seizure"):
                                # Fallback local regex is noisy; prefer Gemini LLM
                                seized_data["total_drug_quantity"] = g_meth
                                seized_data["meth_pills"] = g_meth
                            else:
                                seized_data["total_drug_quantity"] = g_meth
                                seized_data["meth_pills"] = g_meth
                        if g_seized.get("ice_grams", 0) > 0 and (has_vision or seized_data.get("ice_grams", 0) == 0):
                            seized_data["ice_grams"] = float(g_seized["ice_grams"])
                        if g_seized.get("firearms", 0) > 0 and (has_vision or seized_data.get("weapons_count", 0) == 0):
                            seized_data["weapons_count"] = int(g_seized["firearms"])
                        if g_seized.get("ammunition", 0) > 0 and (has_vision or seized_data.get("ammunition_count", 0) == 0):
                            seized_data["ammunition_count"] = int(g_seized["ammunition"])
                        if g_seized.get("cash_thb", 0) > 0 and (has_vision or seized_data.get("cash_amount", 0) == 0):
                            seized_data["cash_amount"] = float(g_seized["cash_thb"])
                        
                        # Merge vehicles
                        for v in g_seized.get("vehicles", []):
                            if v:
                                # Validate plate syntax using strict standard function
                                v_valid_list = extract_vehicle_plates(v)
                                for valid_v in v_valid_list:
                                    if valid_v not in result.get("vehicle_plates", []):
                                        result.setdefault("vehicle_plates", []).append(valid_v)

                    # Update drug_quantity across suspects if enriched
                    if seized_data.get("total_drug_quantity", 0) > 0 and result.get("suspects"):
                        for s in result["suspects"]:
                            if has_vision or not s.get("drug_quantity") or s.get("drug_quantity") == 0:
                                s["drug_quantity"] = seized_data["total_drug_quantity"]

                    # Enrich Scene Location (isolate from arresting station)
                    g_scene = gemini_data.get("crime_scene", {})
                    if g_scene:
                        if not result.get("district") or (result.get("district") == "นางรอง" and g_scene.get("district") and g_scene.get("district") != "นางรอง"):
                            if g_scene.get("district"):
                                result["district"] = g_scene["district"]
                        if g_scene.get("subdistrict") and not result.get("subdistrict"):
                            result["subdistrict"] = g_scene["subdistrict"]
                        if g_scene.get("location") and (has_vision or not result.get("incident_location")):
                            result["incident_location"] = g_scene["location"]

                    # Enrich Case Number & Behavior
                    if gemini_data.get("case_number") and not result.get("case_number"):
                        result["case_number"] = gemini_data["case_number"]
                    if gemini_data.get("behavior") and (has_vision or not result.get("behavior")):
                        result["behavior"] = gemini_data["behavior"]

                    logger.info(f"Gemini Quality Check enriched document: {filename} (Suspects: {len(result['suspects'])})")
            except Exception as e:
                logger.warning(f"Gemini Quality Check encountered error: {e}")

        # If behavior is not yet set but we have seized items or case info, construct meaningful summary
        if not result.get("behavior") and (seized_data.get("summary") or plates_data or social_data):
            summary_parts = []
            if case_data.get("case_number"):
                summary_parts.append(f"คดีอาญาที่ {case_data['case_number']}")
            if seized_data.get("summary"):
                summary_parts.append("ของกลาง: " + ", ".join(seized_data["summary"]))
            if plates_data:
                summary_parts.append("ยานพาหนะ: " + ", ".join(plates_data))
            if social_data:
                summary_parts.append("โซเชียล: " + ", ".join([f"{sm['platform']}: {sm['account']}" for sm in social_data]))
            result["behavior"] = " | ".join(summary_parts)

        return result
        
    except Exception as e:
        logger.error(f"Document extraction failed for {filename}: {e}", exc_info=True)
        return {"suspects": [], "behavior": None}

